"""
Import presets through start_import_session (#34000): an export recognised by a preset is cleaned
into one file per target, and each opens an ordinary import session with the mapping applied; a
file no preset recognises opens the normal way, in the same call. These tests use a preset patched
in; the shipped presets are tested at the end of this file.
"""

import csv
import io
import json
import os
from unittest.mock import patch

import frappe
from openpyxl import Workbook
from openpyxl.styles import Alignment

from frappe_assistant_core.core.tool_registry import get_tool_registry
from frappe_assistant_core.plugins.migration.tools import start_import_session as tool_module
from frappe_assistant_core.tests.base_test import BaseAssistantTest
from frappe_assistant_core.tests.test_import_opening_balances import OpeningBalancesCase
from frappe_assistant_core.tests.test_import_presets import LEDGERS, TB
from frappe_assistant_core.utils import import_opening_invoices as opening_invoices
from frappe_assistant_core.utils import import_opening_stock as opening_stock
from frappe_assistant_core.utils import import_presets as presets
from frappe_assistant_core.utils.import_mapping import build_import_schema
from frappe_assistant_core.utils.import_session import import_target_problem


def xlsx(rows, sheet="Sheet1") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def has_contacts_table():
    """Whether Customer's import takes contacts in a table of their own (Frappe 16's provider)."""
    schema = build_import_schema("Customer")
    return any(t["fieldname"] == "contacts" for t in schema.get("child_tables", []))


def mapped_columns(described):
    """Column → field of a described mapping, the record's and its tables' (contacts.email_id)."""
    rows = [
        *described["mapping"],
        *(row for table in described.get("child_tables", {}).values() for row in table),
    ]
    return {row["column"]: row["field"] for row in rows}


def assert_gstin_and_email_follow_the_site(test, mapped, result):
    """The shipped ledgers preset's GSTIN and Email columns, as this site's Customer takes them.

    GSTIN goes to the first of ["gstin", "tax_id"] the site has: India Compliance adds gstin.
    Email goes to the Customer's email_id; on ERPNext 16, whose import provider takes a party's
    contacts in a table of their own, to contacts.email_id. Where an administrator hid email_id
    (Frappe 15), it is reported as not on this site.
    """
    customer = frappe.get_meta("Customer")
    test.assertEqual(mapped["GSTIN"], "gstin" if customer.has_field("gstin") else "tax_id")
    if has_contacts_table():
        test.assertEqual(mapped["Email"], "contacts.email_id")
        test.assertNotIn("Email (email_id)", result.get("not_on_this_site", []))
    elif customer.get_field("email_id").hidden:
        test.assertNotIn("Email", mapped)
        test.assertIn("Email (email_id)", result["not_on_this_site"])
    else:
        test.assertEqual(mapped["Email"], "email_id")
        test.assertNotIn("Email (email_id)", result.get("not_on_this_site", []))


LEDGER_ROWS = [
    ["Acme Pvt Ltd"],
    ["Ledgers"],
    [],
    ["Name", "Under", "GSTIN", "Opening Balance"],
    ["Ravi Traders {tag}", "Sundry Debtors", "29ABCDE1234F1Z5", "1,000 Dr"],
    ["Mehta Supplies {tag}", "Sundry Creditors", "", "500 Cr"],
    ["HDFC Bank {tag}", "Bank Accounts", "", ""],
    ["Grand Total", "", "", "500"],
]


class PresetToolCase(BaseAssistantTest):
    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        self.registry = get_tool_registry()
        self.tag = frappe.generate_hash(length=6)
        self.files_before = set(frappe.get_all("File", pluck="name"))
        presets.clear_cache()

    def tearDown(self):
        frappe.set_user("Administrator")
        presets.clear_cache()
        for name in set(frappe.get_all("File", pluck="name")) - self.files_before:
            frappe.delete_doc("File", name, force=True, ignore_permissions=True)
        super().tearDown()

    def _upload(self, file_name, content):
        return frappe.get_doc(
            {"doctype": "File", "file_name": file_name, "content": content, "is_private": 1}
        ).insert()

    def _rows(self, rows):
        return [[str(c).format(tag=self.tag) for c in row] for row in rows]

    def _run(self, arguments, use=(LEDGERS,)):
        with patch.object(presets, "load_presets", return_value=list(use)):
            return self.registry.get_tool("start_import_session")._safe_execute(arguments)

    def _apply(self, arguments, use=(LEDGERS,)):
        reply = self._run(arguments, use)
        self.assertTrue(reply["success"], reply)
        return reply["result"]


class TestApply(PresetToolCase):
    def test_a_recognised_file_opens_one_mapped_session_per_target(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        result = self._apply({"file_id": file_doc.name})
        self.assertTrue(result["matched"])
        self.assertEqual(result["preset"]["id"], "test.ledgers")
        self.assertIn("Test", result["message"])
        self.assertEqual(result["source"]["header_row"], 4)
        self.assertEqual(result["dropped"], {"title": 3, "totals": 1})
        by_target = {s["target"]: s for s in result["sessions"]}
        self.assertEqual(set(by_target), {"Customers", "Suppliers"})
        self.assertEqual(
            result["needs"],
            [{"target": "Accounts", "rows": 1, "name": "company", "ask": "Which company?"}],
        )

        customers = by_target["Customers"]
        session = frappe.get_doc("FAC Import Session", customers["session_id"])
        self.assertEqual(session.status, "Mapped")
        self.assertEqual(session.target_doctype, "Customer")
        options = json.loads(session.template_options)
        self.assertEqual(options["column_to_field_map"]["0"], "customer_name")
        self.assertEqual(options["column_to_field_map"]["4"], "customer_type")  # the constant column
        cleaned = frappe.get_doc("File", session.source_file)
        self.assertEqual(cleaned.attached_to_doctype, "FAC Import Session")
        self.assertEqual(cleaned.attached_to_name, session.name)
        content = cleaned.get_content()
        self.assertEqual(
            content.decode() if isinstance(content, bytes) else content,
            f"Name,Group,GSTIN,Opening,Customer Type\nRavi Traders {self.tag},Sundry Debtors,29ABCDE1234F1Z5,1000,Company\n",
        )
        self.assertEqual([s.step for s in session.steps], ["Read File", "Apply Preset", "Map Columns"])
        self.assertEqual(customers["sample_rows"][0][0], f"Ravi Traders {self.tag}")
        self.assertEqual(customers["next"], "dry_run_import")

    def test_a_second_call_with_the_input_opens_only_the_named_target(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        self._apply({"file_id": file_doc.name})
        before = frappe.db.count("FAC Import Session")
        company = frappe.get_all("Company", pluck="name", limit=1)[0]
        result = self._apply(
            {"file_id": file_doc.name, "values": {"company": company}, "targets": ["Accounts"]}
        )
        if import_target_problem("Account"):
            # Account can't be imported on this site: the target is skipped with the reason.
            self.assertEqual(result["sessions"], [])
            self.assertEqual([s["target"] for s in result["skipped_targets"]], ["Accounts"])
            self.assertEqual(frappe.db.count("FAC Import Session"), before)
        else:
            self.assertEqual([s["target"] for s in result["sessions"]], ["Accounts"])
            self.assertEqual(frappe.db.count("FAC Import Session"), before + 1)
        self.assertNotIn("needs", result)

    def test_gstin_goes_to_the_first_field_this_site_has(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        preset = json.loads(json.dumps(LEDGERS))
        preset["targets"][0]["mapping"]["GSTIN"] = ["gstin_not_here", "tax_id"]
        preset["targets"][0]["mapping"]["Opening"] = ["pan_not_here"]
        result = self._apply({"file_id": file_doc.name}, use=(preset,))
        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        session = frappe.get_doc("FAC Import Session", customers["session_id"])
        options = json.loads(session.template_options)
        self.assertEqual(options["column_to_field_map"]["2"], "tax_id")
        self.assertEqual(options["column_to_field_map"]["3"], "Don't Import")
        self.assertIn("Opening (pan_not_here)", result["not_on_this_site"])

    def test_an_unverified_preset_carries_a_warning(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        result = self._apply({"file_id": file_doc.name})
        self.assertFalse(result["preset"]["verified"])
        self.assertTrue(any("not checked against a real" in w for w in result["warnings"]))

    def test_an_engine_target_returns_columns_for_its_setup_tool(self):
        tb = json.loads(json.dumps(TB))
        rows = [
            ["Acme"],
            ["Trial Balance"],
            ["Particulars", "Debit", "Credit"],
            ["Cash", "100", ""],
            ["Capital", "", "100"],
            ["Grand Total", "100", "100"],
        ]
        file_doc = self._upload(f"tb-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        result = self._apply({"file_id": file_doc.name}, use=(tb,))
        session = result["sessions"][0]
        self.assertEqual(session["engine"], "opening_balances")
        self.assertEqual(session["columns"], {"Account": "account", "Debit": "debit", "Credit": "credit"})
        self.assertEqual(session["next"], "set_opening_balances")
        doc = frappe.get_doc("FAC Import Session", session["session_id"])
        self.assertEqual(doc.status, "File Read")
        self.assertIsNone(doc.target_doctype)

    def test_no_match_opens_the_file_the_normal_way(self):
        # One call per file: a file no preset recognises gets its session from the same call.
        file_doc = self._upload(f"plain-{self.tag}.csv", b"Customer Name,Customer Group\nA,Retail\n")
        before = frappe.db.count("FAC Import Session")
        result = self._apply({"file_id": file_doc.name})
        self.assertNotIn("matched", result)
        self.assertEqual(result["columns"], ["Customer Name", "Customer Group"])
        self.assertEqual(result["row_count"], 1)
        self.assertEqual(frappe.db.count("FAC Import Session"), before + 1)
        session = frappe.get_doc("FAC Import Session", result["session_id"])
        self.assertEqual(session.status, "File Read")
        self.assertEqual([s.step for s in session.steps], ["Read File"])

    def test_preset_none_opens_a_recognised_export_the_normal_way(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        result = self._apply({"file_id": file_doc.name, "preset": "none"})
        self.assertNotIn("matched", result)
        self.assertEqual(result["columns"], ["Name", "Under", "GSTIN", "Opening Balance"])
        self.assertEqual(result["header_row"], 4)

    def test_a_heading_row_given_by_the_user_opens_the_file_the_normal_way(self):
        # header_row is the user's answer about this file's headings: no preset second-guesses it.
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        result = self._apply({"file_id": file_doc.name, "header_row": 4})
        self.assertNotIn("matched", result)
        self.assertEqual(result["columns"], ["Name", "Under", "GSTIN", "Opening Balance"])

    def test_a_google_sheet_is_never_matched_against_a_preset(self):
        for arguments in ({"preset": "test.ledgers"}, {"parts": ["ledgers"]}):
            reply = self._run({"google_sheets_url": "https://docs.google.com/spreadsheets/d/x", **arguments})
            self.assertFalse(reply["success"], arguments)
            self.assertIn("a file on this site", reply["error"])

    def test_the_preset_tool_is_gone_and_start_import_session_takes_its_inputs(self):
        self.assertIsNone(self.registry.get_tool("apply_import_preset"))
        inputs = self.registry.get_tool("start_import_session").inputSchema["properties"]
        for name in ("preset", "values", "targets", "parts"):
            self.assertIn(name, inputs)

    def test_a_named_preset_that_does_not_fit_says_which_headers_are_missing(self):
        file_doc = self._upload(f"plain-{self.tag}.csv", b"Name,GSTIN\nA,B\n")
        result = self._apply({"file_id": file_doc.name, "preset": "test.ledgers"})
        self.assertFalse(result["matched"])
        self.assertEqual(result["missing_headers"], [["Under", "Group Name"]])

    def test_a_named_preset_without_its_title_says_so(self):
        rows = [["Particulars", "Debit", "Credit"], ["Cash", "100", ""]]
        file_doc = self._upload(f"tb-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        result = self._apply({"file_id": file_doc.name, "preset": "test.tb"}, use=(TB,))
        self.assertFalse(result["matched"])
        self.assertEqual(result["missing_headers"], [])
        self.assertEqual(result["missing_title"], ["*trial balance*"])

    def test_a_named_preset_short_of_optional_groups_says_how_many(self):
        strict = {**LEDGERS, "signature": {**LEDGERS["signature"], "min_optional": 1}}
        file_doc = self._upload(f"plain-{self.tag}.csv", b"Name,Under\nA,B\n")
        result = self._apply({"file_id": file_doc.name, "preset": "test.ledgers"}, use=(strict,))
        self.assertFalse(result["matched"])
        self.assertEqual(result["missing_headers"], [])
        self.assertEqual(result["needs_optional"], {"min": 1, "groups": [["GSTIN"], ["Missing"]]})

    def test_an_unknown_preset_is_refused_with_the_list(self):
        file_doc = self._upload(f"plain-{self.tag}.csv", b"Name,Under\nA,B\n")
        reply = self._run({"file_id": file_doc.name, "preset": "test.nothing"})
        self.assertFalse(reply["success"])
        self.assertIn("test.ledgers", reply["error"])

    def test_only_title_and_totals_left_is_an_error(self):
        rows = [["Title"], ["Name", "Under"], ["Grand Total", ""]]
        file_doc = self._upload(f"empty-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        reply = self._run({"file_id": file_doc.name})
        self.assertFalse(reply["success"])
        self.assertIn("no data rows", reply["error"])
        self.assertTrue(
            reply["error"].endswith(
                "If this is not a Test export, call start_import_session again with preset 'none'."
            ),
            reply["error"],
        )

    def test_a_file_no_preset_matches_is_answered_from_its_head(self):
        rows = [["Customer Name", "Customer Group"]] + [[f"C{i}", "Retail"] for i in range(300)]
        file_doc = self._upload(f"plain-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        with patch.object(presets, "read_sheet", side_effect=AssertionError("read the whole sheet")):
            result = self._apply({"file_id": file_doc.name})
        self.assertNotIn("matched", result)
        self.assertEqual(result["file_name"], f"plain-{self.tag}.csv")
        self.assertEqual(result["row_count"], 300)

    def test_only_the_matched_sheet_is_read_in_full(self):
        workbook_rows = {"Notes": [["Nothing here"]], "Ledgers": self._rows(LEDGER_ROWS)}
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, rows in workbook_rows.items():
            sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        file_doc = self._upload(f"two-{self.tag}.xlsx", buffer.getvalue())
        with patch.object(presets, "read_sheet", wraps=presets.read_sheet) as read_sheet:
            result = self._apply({"file_id": file_doc.name})
        self.assertEqual(read_sheet.call_count, 1)
        self.assertEqual(read_sheet.call_args.args[1], "Ledgers")
        self.assertEqual(result["source"]["sheet"], "Ledgers")
        self.assertEqual(result["source"]["header_row"], 4)

    def test_a_matched_sheet_over_the_row_ceiling_is_refused(self):
        rows = [["Name", "Under"]] + [[f"P{i}", "Sundry Debtors"] for i in range(5)]
        file_doc = self._upload(f"big-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        before = frappe.db.count("FAC Import Session")
        with patch.object(presets, "MAX_PRESET_ROWS", 5):
            reply = self._run({"file_id": file_doc.name})
        self.assertFalse(reply["success"])
        self.assertIn(
            "This sheet has more than 5 rows. Split it into smaller files, or call start_import_session "
            "again with preset 'none'.",
            reply["error"],
        )
        self.assertEqual(frappe.db.count("FAC Import Session"), before)

    def test_a_sheet_given_for_a_csv_is_ignored(self):
        rows = [["Name", "Under"], [f"Ravi {self.tag}", "Sundry Debtors"]]
        file_doc = self._upload(f"ledgers-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        result = self._apply({"file_id": file_doc.name, "sheet": "Sheet1"})
        self.assertTrue(result["matched"])
        self.assertEqual([s["target"] for s in result["sessions"]], ["Customers"])

    def test_a_sheet_that_is_not_in_the_workbook_is_refused(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        reply = self._run({"file_id": file_doc.name, "sheet": "Nope"})
        self.assertFalse(reply["success"])
        self.assertIn("Sheet 'Nope' not found.", reply["error"])

    def test_values_and_targets_of_the_wrong_type_are_refused(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        before = frappe.db.count("FAC Import Session"), frappe.db.count("File")
        tool = self.registry.get_tool("start_import_session")
        # The tool's own check. Through _safe_execute the schema check refuses a value that is not
        # an object or a list before execute runs, so execute is called directly here.
        for arguments, error in (
            ({"values": "Acme Ltd"}, 'values must be an object like {"company": "Acme Ltd"}'),
            ({"values": ["Acme Ltd"]}, 'values must be an object like {"company": "Acme Ltd"}'),
            ({"targets": "Customers"}, "targets must be a list of target names"),
            ({"targets": ["Customers", 3]}, "targets must be a list of target names"),
        ):
            with patch.object(presets, "load_presets", return_value=[LEDGERS]):
                reply = tool.execute({"file_id": file_doc.name, **arguments})
            self.assertEqual(reply, {"success": False, "error": error}, arguments)
        # A list with a name that is not a string passes the schema check and is refused by the tool.
        reply = self._run({"file_id": file_doc.name, "targets": ["Customers", 3]})
        self.assertFalse(reply["success"])
        self.assertIn("targets must be a list of target names", reply["error"])
        self.assertEqual((frappe.db.count("FAC Import Session"), frappe.db.count("File")), before)

    def test_a_user_who_cannot_read_the_file_gets_the_usual_refusal(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        outsider = self.make_desk_user("outsider")
        self.assert_cannot(outsider, "read", "File", file_doc.name)
        frappe.set_user(outsider)
        before = frappe.db.count("FAC Import Session")
        reply = self._run({"file_id": file_doc.name})
        self.assertFalse(reply["success"])
        self.assertIn("File not found, or you don't have access to it.", reply["error"])
        frappe.set_user("Administrator")
        self.assertEqual(frappe.db.count("FAC Import Session"), before)

    def test_a_wrong_file_type_gets_the_usual_refusal(self):
        file_doc = self._upload(f"notes-{self.tag}.txt", b"Name,Under\n")
        before = frappe.db.count("FAC Import Session"), frappe.db.count("File")
        reply = self._run({"file_id": file_doc.name})
        self.assertFalse(reply["success"])
        self.assertIn("Only .csv, .xlsx or .xls", reply["error"])
        self.assertEqual((frappe.db.count("FAC Import Session"), frappe.db.count("File")), before)

    def test_a_file_over_50_mb_gets_the_usual_refusal_and_writes_nothing(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        # The size Frappe recorded at upload is what both tools check; 50 MB is not written.
        frappe.db.set_value("File", file_doc.name, "file_size", 51 * 1024 * 1024)
        before = frappe.db.count("FAC Import Session"), frappe.db.count("File")
        reply = self._run({"file_id": file_doc.name})
        self.assertFalse(reply["success"])
        self.assertIn("The file is larger than", reply["error"])  # the site's Max File Size (#34413)
        self.assertEqual((frappe.db.count("FAC Import Session"), frappe.db.count("File")), before)

    def test_a_target_the_user_cannot_import_is_skipped_with_the_reason(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))

        def refuse_supplier(doctype):
            return "no Supplier for you" if doctype == "Supplier" else None

        with patch.object(tool_module, "import_target_problem", side_effect=refuse_supplier):
            result = self._apply({"file_id": file_doc.name})
        self.assertEqual([s["target"] for s in result["sessions"]], ["Customers"])
        self.assertEqual(
            result["skipped_targets"],
            [{"target": "Suppliers", "rows": 1, "reason": "no Supplier for you"}],
        )

    def test_set_column_mapping_still_changes_a_column_afterwards(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        session = next(
            s for s in self._apply({"file_id": file_doc.name})["sessions"] if s["target"] == "Customers"
        )
        reply = self.registry.get_tool("set_column_mapping")._safe_execute(
            {"session": session["session_id"], "mapping": {"Group": "customer_group"}}
        )
        self.assertTrue(reply["success"], reply)
        options = json.loads(frappe.get_doc("FAC Import Session", session["session_id"]).template_options)
        self.assertEqual(options["column_to_field_map"]["1"], "customer_group")
        self.assertEqual(options["column_to_field_map"]["0"], "customer_name")

    def test_two_columns_with_the_same_fallback_do_not_take_the_same_field(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        preset = json.loads(json.dumps(LEDGERS))
        preset["targets"][0]["mapping"]["GSTIN"] = ["gstin_not_here", "tax_id"]
        preset["targets"][0]["mapping"]["Opening"] = ["gstin_not_here", "tax_id"]
        result = self._apply({"file_id": file_doc.name}, use=(preset,))
        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        self.assertNotIn("mapping_problems", customers)
        self.assertEqual(customers["next"], "dry_run_import")
        options = json.loads(frappe.get_doc("FAC Import Session", customers["session_id"]).template_options)
        self.assertEqual(options["column_to_field_map"]["2"], "tax_id")
        self.assertEqual(options["column_to_field_map"]["3"], "Don't Import")
        # tax_id is on this site, taken by GSTIN: Opening is already mapped, not missing here.
        self.assertEqual(result["already_mapped"], ["Opening (gstin_not_here / tax_id)"])
        self.assertNotIn("Opening (gstin_not_here / tax_id)", result.get("not_on_this_site", []))

    def test_a_mapping_that_cannot_be_saved_points_to_set_column_mapping(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        with (
            patch.object(tool_module, "resolve_mapping", return_value=({}, ["Column X has no field"])),
            patch.object(tool_module, "duplicate_targets", return_value=[]),
        ):
            result = self._apply({"file_id": file_doc.name})
        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        self.assertEqual(customers["mapping_problems"], ["Column X has no field"])
        self.assertNotIn("mapping", customers)
        self.assertEqual(customers["next"], "set_column_mapping")

    def test_a_failure_opening_one_target_leaves_the_others_and_no_half_writes(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        private_files = frappe.get_site_path("private", "files")
        real_save = tool_module.save_mapping

        def save_unless_supplier(session, doctype, *args, **kwargs):
            if doctype == "Supplier":
                raise ValueError("cannot save the Supplier mapping")
            return real_save(session, doctype, *args, **kwargs)

        with patch.object(tool_module, "save_mapping", side_effect=save_unless_supplier):
            result = self._apply({"file_id": file_doc.name})
        self.assertEqual([s["target"] for s in result["sessions"]], ["Customers"])
        self.assertEqual(
            result["skipped_targets"],
            [{"target": "Suppliers", "rows": 1, "reason": "cannot save the Supplier mapping"}],
        )
        # This test's sessions only: a site used for live tests has Supplier sessions of its own
        self.assertEqual(
            frappe.db.count(
                "FAC Import Session",
                {"target_doctype": "Supplier", "file_name": ["like", f"ledgers-{self.tag} - %"]},
            ),
            0,
        )
        self.assertEqual(
            frappe.db.count("File", {"file_name": ["like", f"ledgers-{self.tag} - Suppliers%"]}), 0
        )
        self.assertEqual(
            frappe.db.count("File", {"file_name": ["like", f"ledgers-{self.tag} - Customers%"]}), 1
        )
        # The cleaned Suppliers file written to disk before the failure is gone too.
        on_disk = [f for f in os.listdir(private_files) if f.startswith(f"ledgers-{self.tag} - ")]
        self.assertTrue(any("Customers" in f for f in on_disk), on_disk)
        self.assertFalse(any("Suppliers" in f for f in on_disk), on_disk)

    def test_an_engine_doctype_missing_from_the_site_is_skipped(self):
        rows = [
            ["Acme"],
            ["Trial Balance"],
            ["Particulars", "Debit", "Credit"],
            ["Cash", "100", ""],
            ["Capital", "", "100"],
        ]
        file_doc = self._upload(f"tb-{self.tag}.csv", "\n".join(",".join(r) for r in rows).encode())
        with patch.object(presets, "engine_doctype", return_value="No Such Doctype"):
            result = self._apply({"file_id": file_doc.name}, use=(TB,))
        self.assertEqual(result["sessions"], [])
        self.assertEqual(
            result["skipped_targets"][0]["reason"], "No Such Doctype isn't available on this site."
        )

    def test_an_unknown_target_name_is_refused_and_writes_nothing(self):
        file_doc = self._upload(f"ledgers-{self.tag}.xlsx", xlsx(self._rows(LEDGER_ROWS)))
        before = frappe.db.count("FAC Import Session"), frappe.db.count("File")
        reply = self._run({"file_id": file_doc.name, "targets": ["Nope"]})
        self.assertFalse(reply["success"])
        self.assertIn(
            "No target 'Nope'. Targets of test.ledgers: Customers, Suppliers, Accounts.", reply["error"]
        )
        self.assertEqual((frappe.db.count("FAC Import Session"), frappe.db.count("File")), before)


# Each fixture follows the layout in the spec's section 7 / the research notes: (rows, {target: rows}).
SHIPPED = {
    "tally_prime.ledgers": (
        [
            ["Acme Pvt Ltd"],
            ["List of Ledgers"],
            [],
            ["Name", "Group Name", "GST Registration -GSTIN/UIN", "E-mail", "Ledger – Opening Balance"],
            ["Ravi Traders", "Sundry Debtors", "29ABCDE1234F1Z5", "ravi@example.com", "1,000.00 Dr"],
            ["Mehta Supplies", "Sundry Creditors", "", "", "500.00 Cr"],
            ["Rent", "Indirect Expenses", "", "", ""],
        ],
        {"Customers": 1, "Suppliers": 1},  # Accounts waits for the company (needs)
    ),
    "tally_prime.stock_items": (
        [
            ["Acme Pvt Ltd"],
            [],
            ["Name", "Group Name", "Units", "HSN/SAC", "Opening Balance – Quantity"],
            ["Steel Rod 10mm", "Raw Material", "Nos", "7214", "10 Nos"],
            ["Paint 1L", "Finished Goods", "Ltr", "3208", ""],
        ],
        {"Items": 2},
    ),
    "tally_prime.trial_balance": (
        [
            ["Acme Pvt Ltd"],
            ["1-Apr-24 to 31-Mar-25"],
            [],
            ["Particulars", "Debit", "Credit"],
            ["Current Assets", "1,500.00", ""],
            ["  Cash", "500.00", ""],
            ["  HDFC Bank", "1,000.00", ""],
            ["Capital Account", "", "1,500.00"],
            ["  Owner Capital", "", "1,500.00"],
            ["Grand Total", "1,500.00", "1,500.00"],
        ],
        {"Trial balance": 3},
    ),
    "tally_prime.stock_summary": (
        [
            ["Acme Pvt Ltd"],
            ["Stock Summary"],
            ["1-Apr-24 to 31-Mar-25"],
            ["Particulars", "Quantity", "Rate", "Value"],
            ["Raw Material", "", "", "1,200.00"],
            ["  Steel Rod 10mm", "10 Nos", "100.00/Nos", "1,000.00"],
            ["  Paint 1L", "4 Ltr", "50.00/Ltr", "200.00"],
            ["Grand Total", "", "", "1,200.00"],
        ],
        {},  # needs a warehouse first
    ),
    "tally_prime.bills_receivable": (
        [
            ["Acme Pvt Ltd"],
            ["Bills Receivable"],
            ["1-Apr-24 to 31-Mar-25"],
            ["Date", "Ref. No.", "Party's Name", "Pending Amount", "Due on"],
            ["1-Mar-25", "S-101", "Ravi Traders", "1,000.00 Dr", "31-Mar-25"],
            ["Grand Total", "", "", "1,000.00 Dr", ""],
        ],
        {"Open invoices": 1},
    ),
    "tally_prime.bills_payable": (
        [
            ["Acme Pvt Ltd"],
            ["Bills Payable"],
            ["1-Apr-24 to 31-Mar-25"],
            ["Date", "Ref. No.", "Party's Name", "Pending Amount", "Due on"],
            ["2-Mar-25", "P-77", "Mehta Supplies", "500.00 Cr", "1-Apr-25"],
            ["3-Mar-25", "P-78", "Mehta Supplies", "250.00", "2-Apr-25"],
        ],
        {"Open bills": 2},
    ),
    "quickbooks_online.account_list": (
        [
            ["Account List"],
            ["Sandbox Company_US_2"],
            [],
            ["Full name", "Type", "Detail type", "Description", "Total balance"],
            ["Checking", "Bank", "Checking", "", "1,201.00"],
            ["Truck", "Fixed Assets", "Vehicles", "", ""],
            ["Truck:Depreciation", "Fixed Assets", "Accumulated Depreciation", "", "-500.00"],
            ["TOTAL", "", "", "", "$701.00"],
            [],
            [],
            [],
            [" Wednesday, March 04, 2026 10:53 PM GMTZ"],
        ],
        {},  # needs the company first
    ),
    "quickbooks_online.trial_balance": (
        [
            ["Trial Balance"],
            ["Sandbox Company_US_2"],
            ["As of December 31, 2025"],
            [],
            ["", "Debit", "Credit"],
            ["Checking", "1,201.00", ""],
            ["Opening Balance Equity", "", "1,201.00"],
            ["TOTAL", "$1,201.00", "$1,201.00"],
            [],
            [],
            [],
            ["Accrual Basis Wednesday, March 04, 2026 11:09 PM GMTZ"],
        ],
        {"Trial balance": 2},
    ),
    "quickbooks_online.ar_aging_summary": (
        [
            ["A/R Aging Summary Report"],
            ["Sandbox Company_US_2"],
            ["As of March 4, 2026"],
            [],
            ["Customer", "CURRENT", "1 - 30", "31 - 60", "61 - 90", "91 AND OVER", "Total"],
            ["Amy's Bird Sanctuary", "", "", "239.00", "", "", "239.00"],
            ["Freeman Sporting Goods", "", "", "", "", "", ""],
            ["0969 Ocean View Road", "", "", "477.50", "", "", "477.50"],
            ["Total for Freeman Sporting Goods", "", "", "477.50", "", "", "$477.50"],
            ["TOTAL", "", "", "716.50", "", "", "$716.50"],
        ],
        {"Open invoices": 2},
    ),
    "quickbooks_online.ap_aging_summary": (
        [
            ["A/P Aging Summary Report"],
            ["Sandbox Company_US_2"],
            ["As of March 4, 2026"],
            [],
            ["Vendor", "CURRENT", "1 - 30", "31 - 60", "61 - 90", "91 AND OVER", "Total"],
            ["Brosnahan Insurance Agency", "241.23", "", "", "", "", "241.23"],
        ],
        {"Open bills": 1},
    ),
    "quickbooks_online.customer_contact_list": (
        [
            ["Customer Contact List"],
            ["Sandbox Company_US_2"],
            ["All Dates"],
            [],
            ["Customer", "Phone Numbers", "Email", "Full Name", "Billing Address", "Shipping Address"],
            [
                "Test Contact Echo",
                "Phone: 555-010-0011",
                "echo.test@example.com",
                "Echo Tester",
                "400 Mock Lane",
                "",
            ],
        ],
        {"Customers": 1},
    ),
    "quickbooks_online.vendor_contact_list": (
        [
            ["Vendor Contact List"],
            ["Sandbox Company_US_2"],
            ["All Dates"],
            [],
            ["Vendor", "Phone Numbers", "Email", "Full Name", "Address", "Account #"],
            ["Brosnahan Insurance Agency", "Phone: (650) 555-9912", "", "Nick Brosnahan", "P.O. Box 5", ""],
        ],
        {"Suppliers": 1},
    ),
    "quickbooks_online.products_services": (
        [
            ["Product/Service", "Type", "Description", "Price", "Cost", "Qty on hand"],
            ["Garden Rocks", "Inventory", "Decorative rocks", "35.00", "20.00", "12"],
            ["Design", "Service", "Custom design", "75.00", "", ""],
        ],
        {"Items": 2},
    ),
}


SHIPPED.update(
    {
        "quickbooks_desktop.customer_list": (
            [
                ["Active Status", "Customer", "Balance", "Company", "Main Phone", "Main Email", "Bill to 1"],
                [
                    "Active",
                    "Abercrombie, Kristy",
                    "0.00",
                    "",
                    "415-555-6579",
                    "kristy@example.com",
                    "Kristy Abercrombie",
                ],
                ["Inactive", "Old Co", "0.00", "Old Co", "", "", ""],
            ],
            {"Customers": 2},
        ),
        "quickbooks_desktop.vendor_list": (
            [
                ["Active Status", "Vendor", "Balance", "Company", "Main Phone", "Main Email", "Bill from 1"],
                ["Active", "Cal Telephone", "0.00", "Cal Telephone", "415-555-1234", "", "Cal Telephone"],
            ],
            {"Suppliers": 1},
        ),
        "quickbooks_desktop.item_list": (
            [
                ["Item", "Description", "Type", "Account", "On Hand", "Price", "Cost"],
                [
                    "Cabinets:Cabinet Pulls",
                    "Cabinet Pulls",
                    "Inventory Part",
                    "Construction Income",
                    "423",
                    "0.00",
                    "3.00",
                ],
                ["Framing", "Framing labor", "Service", "Construction Income", "", "55.00", ""],
                ["Subtotal", "", "Subtotal", "", "", "", ""],
            ],
            {"Items": 2},
        ),
        "quickbooks_desktop.account_listing": (
            [
                ["Account", "Type", "Balance Total", "Description", "Accnt. #", "Tax Line"],
                ["1000 · Checking", "Bank", "46,969.10", "", "1000", ""],
                ["1500 · Truck:1510 · Depreciation", "Fixed Asset", "-2,500.00", "", "1510", ""],
                ["Estimates", "Non-Posting", "0.00", "", "", ""],
            ],
            {},  # needs the company first; Non-Posting is no_target
        ),
        "zoho_books.customers": (
            [
                [
                    "Display Name",
                    "Company Name",
                    "EmailID",
                    "Phone",
                    "MobilePhone",
                    "GST Treatment",
                    "GST Identification Number (GSTIN)",
                    "Currency Code",
                    "Status",
                    "Customer Sub Type",
                    "Contact ID",
                ],
                [
                    "Ravi Traders",
                    "Ravi Traders Pvt Ltd",
                    "ravi@example.com",
                    "080-1234",
                    "",
                    "business_gst",
                    "29ABCDE1234F1Z5",
                    "INR",
                    "Active",
                    "business",
                    "4600000000001",
                ],
            ],
            {"Customers": 1},
        ),
        "zoho_books.vendors": (
            [
                [
                    "Display Name",
                    "Company Name",
                    "EmailID",
                    "Phone",
                    "GST Treatment",
                    "GST Identification Number (GSTIN)",
                    "Status",
                    "Vendor ID",
                    "Bank Account Number",
                ],
                [
                    "Mehta Supplies",
                    "Mehta Supplies",
                    "",
                    "",
                    "business_none",
                    "",
                    "Active",
                    "4600000000002",
                    "",
                ],
            ],
            {"Suppliers": 1},
        ),
        "zoho_books.items": (
            [
                [
                    "Item ID",
                    "Item Name",
                    "SKU",
                    "HSN/SAC",
                    "Description",
                    "Rate",
                    "Usage unit",
                    "Product Type",
                    "Status",
                ],
                ["46000001", "Steel Rod 10mm", "SR10", "7214", "", "INR 100.00", "pcs", "goods", "Active"],
                ["46000002", "Design", "", "998391", "", "75", "", "service", "Inactive"],
            ],
            {"Items": 2},
        ),
        "zoho_books.chart_of_accounts": (
            [
                [
                    "Account ID",
                    "Account Name",
                    "Account Code",
                    "Description",
                    "Account Type",
                    "Account Status",
                    "Currency",
                    "Parent Account",
                ],
                ["1", "Petty Cash", "1001", "", "Cash", "Active", "INR", ""],
                ["2", "Office Rent", "6001", "", "Expense", "Active", "INR", "Rent"],
            ],
            {},  # needs the company first
        ),
        "zoho_books.trial_balance": (
            [
                ["Acme Pvt Ltd"],
                ["Trial Balance"],
                ["Basis: Accrual"],
                ["As of 31/03/2025"],
                ["Account", "Account Code", "Net Debit", "Net Credit"],
                ["Assets", "", "", ""],
                ["Current Assets", "", "", ""],
                ["Petty Cash", "1001", "500.00", ""],
                ["Total for Assets", "", "500.00", ""],
                ["Equities", "", "", ""],
                ["Capital", "3001", "", "500.00"],
                ["Total", "", "500.00", "500.00"],
            ],
            {"Trial balance": 2},
        ),
        "zoho_books.invoices": (
            [
                [
                    "Invoice Date",
                    "Invoice ID",
                    "Invoice Number",
                    "Invoice Status",
                    "Customer ID",
                    "Customer Name",
                    "Due Date",
                    "Balance",
                    "Item Name",
                    "Quantity",
                ],
                [
                    "2025-03-01",
                    "1",
                    "INV-001",
                    "Overdue",
                    "9",
                    "Ravi Traders",
                    "2025-03-31",
                    "1180.00",
                    "Steel Rod",
                    "10",
                ],
                [
                    "2025-03-01",
                    "1",
                    "INV-001",
                    "Overdue",
                    "9",
                    "Ravi Traders",
                    "2025-03-31",
                    "1180.00",
                    "Paint",
                    "2",
                ],
                [
                    "2025-03-05",
                    "2",
                    "INV-002",
                    "Paid",
                    "9",
                    "Ravi Traders",
                    "2025-04-04",
                    "0.00",
                    "Paint",
                    "1",
                ],
            ],
            {"Open invoices": 1},
        ),
        "zoho_books.bills": (
            [
                [
                    "Bill Date",
                    "Bill ID",
                    "Bill Number",
                    "Bill Status",
                    "Vendor Name",
                    "Due Date",
                    "Balance",
                    "Item Name",
                ],
                ["2025-03-02", "7", "B-77", "Open", "Mehta Supplies", "2025-04-01", "500.00", "Steel"],
            ],
            {"Open bills": 1},
        ),
        "busy.accounts": (
            [
                ["Acme Pvt Ltd"],
                ["Account List"],
                ["Name", "Group", "Op. Bal.", "Dr/Cr", "GSTIN", "Mobile", "Email"],
                ["Ravi Traders", "Sundry Debtors", "1000", "Dr", "29ABCDE1234F1Z5", "9876543210", ""],
                ["Mehta Supplies", "Sundry Creditors", "500", "Cr", "", "", ""],
                ["Cash", "Cash-in-hand", "200", "Dr", "", "", ""],
            ],
            {"Customers": 1, "Suppliers": 1},
        ),
        "busy.items": (
            [
                ["Name", "Group", "Main Unit", "HSN Code", "Tax Category", "Sale Price"],
                ["Steel Rod 10mm", "Raw Material", "Pcs.", "7214", "GST 18%", "120.00"],
            ],
            {"Items": 1},
        ),
        "busy.trial_balance": (
            [
                ["Acme Pvt Ltd"],
                ["Trial Balance as on 31-03-2025"],
                ["Account Name", "Debit", "Credit"],
                ["Cash", "200.00", ""],
                ["Capital", "", "200.00"],
                ["Total", "200.00", "200.00"],
            ],
            {"Trial balance": 2},
        ),
        "marg.ledgers": (
            [
                ["Acme Pharma"],
                ["All Ledgers"],
                ["Ledger Name", "Station", "Group", "GSTIN", "Mobile", "Balance", "Dr/Cr"],
                [
                    "Ravi Medicals",
                    "Bengaluru",
                    "Sundry Debtors",
                    "29ABCDE1234F1Z5",
                    "9876543210",
                    "1000",
                    "Dr",
                ],
                ["Cipla Distributor", "Mumbai", "Sundry Creditors", "", "", "500", "Cr"],
            ],
            {"Customers": 1, "Suppliers": 1},
        ),
        "marg.items": (
            [
                ["Product Name", "Packing", "Company", "Unit", "HSN", "GST %", "MRP", "Rate"],
                ["Paracetamol 500", "10x10", "Cipla", "Strip", "3004", "12", "25.00", "20.00"],
            ],
            {"Items": 1},
        ),
    }
)


# Tally's real Trial Balance export: a two-row Closing Balance header and Particulars in a column
# with no header.
TALLY_TB_REAL = [
    ["Acme Pvt Ltd", "", ""],
    ["", "1-Apr-24 to 31-Mar-25", ""],
    ["", "Closing Balance", ""],
    ["", "Debit", "Credit"],
    ["Current Assets", "1,500.00", ""],
    ["  Cash", "500.00", ""],
    ["  HDFC Bank", "1,000.00", ""],
    ["Capital Account", "", "1,500.00"],
    ["  Owner Capital", "", "1,500.00"],
    ["Grand Total", "1,500.00", "1,500.00"],
]


def as_xlsx_file(case, name, rows):
    return case._upload(name, xlsx([[str(c) for c in row] for row in rows]))


class TestShippedPresets(PresetToolCase):
    """Each shipped preset recognises its own layout and splits it as expected."""

    def _check(self, preset_id, rows, expected):
        file_doc = as_xlsx_file(self, f"{preset_id}-{self.tag}.xlsx", rows)
        reply = self.registry.get_tool("start_import_session")._safe_execute({"file_id": file_doc.name})
        self.assertTrue(reply["success"], (preset_id, reply))
        result = reply["result"]
        self.assertTrue(result["matched"], (preset_id, result))
        self.assertEqual(result["preset"]["id"], preset_id, result)
        self.assertEqual({s["target"]: s["rows"] for s in result["sessions"]}, expected, (preset_id, result))
        return result, file_doc

    def test_every_shipped_fixture(self):
        for preset_id, (rows, expected) in SHIPPED.items():
            with self.subTest(preset_id):
                presets.clear_cache()
                result, _ = self._check(preset_id, rows, expected)
                # AC1: a doctype session is mapped, not only counted.
                for session in result["sessions"]:
                    if session.get("doctype"):
                        self.assertIn("mapping", session, (preset_id, session))
                        self.assertNotIn("mapping_problems", session, (preset_id, session))

    def test_every_shipped_preset_has_a_fixture(self):
        self.assertEqual(
            sorted(p["id"] for p in presets.load_presets()),
            sorted(SHIPPED),
        )

    def test_tally_ledgers_ac1(self):
        rows, expected = SHIPPED["tally_prime.ledgers"]
        result, _ = self._check("tally_prime.ledgers", rows, expected)
        self.assertIn("Tally Prime", result["message"])
        self.assertFalse(result["preset"]["verified"])
        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        mapped = mapped_columns(customers["mapping"])
        self.assertEqual(mapped["Name"], "customer_name")
        assert_gstin_and_email_follow_the_site(self, mapped, result)
        self.assertEqual(result["needs"][0]["name"], "company")

    def test_a_hidden_email_field_is_listed_as_not_on_this_site(self):
        if has_contacts_table():
            self.skipTest("Frappe 16 imports the email through the contacts table, hidden or not")
        if not frappe.get_meta("Customer").get_field("email_id").hidden:
            # ERPNext 16 ships it hidden; on v15 an administrator hides it in Customize Form
            from frappe.custom.doctype.property_setter.property_setter import make_property_setter

            setter = make_property_setter("Customer", "email_id", "hidden", 1, "Check")
            self.addCleanup(frappe.clear_cache, doctype="Customer")
            self.addCleanup(frappe.delete_doc, "Property Setter", setter.name, force=True)
        self.assertTrue(frappe.get_meta("Customer").get_field("email_id").hidden)
        rows, expected = SHIPPED["tally_prime.ledgers"]

        result, _ = self._check("tally_prime.ledgers", rows, expected)

        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        mapped = {row["column"]: row["field"] for row in customers["mapping"]["mapping"]}
        self.assertNotIn("Email", mapped)
        self.assertIn("Email", customers["mapping"]["not_imported"])
        self.assertIn("Email (email_id)", result["not_on_this_site"])

    def test_tally_trial_balance_drops_group_subtotals(self):
        rows, expected = SHIPPED["tally_prime.trial_balance"]
        result, _ = self._check("tally_prime.trial_balance", rows, expected)
        self.assertEqual(result["dropped"]["groups"], 2)
        self.assertEqual(
            [r[0] for r in result["sessions"][0]["sample_rows"]], ["Cash", "HDFC Bank", "Owner Capital"]
        )

    def test_tally_trial_balance_with_the_real_two_row_header(self):
        result, _ = self._check("tally_prime.trial_balance", TALLY_TB_REAL, {"Trial balance": 3})
        self.assertEqual(result["source"]["header_row"], 4)  # 1-based: the Debit/Credit row
        self.assertEqual(result["dropped"]["groups"], 2)
        sample = result["sessions"][0]["sample_rows"]
        self.assertEqual(
            [r[:3] for r in sample],
            [["Cash", "500", ""], ["HDFC Bank", "1000", ""], ["Owner Capital", "", "1500"]],
        )
        self.assertNotIn("Debit", [r[0] for r in sample])

    def test_tally_trial_balance_groups_found_by_xlsx_cell_indent(self):
        # Group names Tally doesn't ship, and ledgers indented by the cell's alignment, not spaces:
        # only read_indents can tell the groups apart.
        rows = [
            (["Acme Pvt Ltd"], 0),
            (["1-Apr-24 to 31-Mar-25"], 0),
            (["Particulars", "Debit", "Credit"], 0),
            (["Office Group", "1,500.00", ""], 0),
            (["Cash", "500.00", ""], 1),
            (["HDFC Bank", "1,000.00", ""], 1),
            (["Owner Group", "", "1,500.00"], 0),
            (["Owner Capital", "", "1,500.00"], 1),
            (["Grand Total", "1,500.00", "1,500.00"], 0),
        ]
        workbook = Workbook()
        worksheet = workbook.active
        for values, indent in rows:
            worksheet.append(values)
            if indent:
                worksheet.cell(row=worksheet.max_row, column=1).alignment = Alignment(indent=indent)
        buffer = io.BytesIO()
        workbook.save(buffer)
        file_doc = self._upload(f"tb-indent-{self.tag}.xlsx", buffer.getvalue())
        reply = self.registry.get_tool("start_import_session")._safe_execute({"file_id": file_doc.name})
        self.assertTrue(reply["success"], reply)
        result = reply["result"]
        self.assertEqual(result["preset"]["id"], "tally_prime.trial_balance")
        self.assertEqual(result["dropped"]["groups"], 2)
        self.assertEqual(
            [r[0] for r in result["sessions"][0]["sample_rows"]], ["Cash", "HDFC Bank", "Owner Capital"]
        )

    def test_a_tally_trial_balance_is_not_claimed_by_quickbooks(self):
        rows = [
            ["Acme Pvt Ltd"],
            ["Trial Balance"],
            ["1-Apr-24 to 31-Mar-25"],
            *SHIPPED["tally_prime.trial_balance"][0][3:],
        ]
        self._check("tally_prime.trial_balance", rows, {"Trial balance": 3})

    def test_tally_bills_payable_reads_dr_cr_and_plain_amounts(self):
        rows, expected = SHIPPED["tally_prime.bills_payable"]
        result, _ = self._check("tally_prime.bills_payable", rows, expected)
        self.assertEqual([r[3] for r in result["sessions"][0]["sample_rows"]], ["500", "250"])

    def test_quickbooks_aging_parties_join_parent_and_child(self):
        rows, expected = SHIPPED["quickbooks_online.ar_aging_summary"]
        result, _ = self._check("quickbooks_online.ar_aging_summary", rows, expected)
        parties = [r[2] for r in result["sessions"][0]["sample_rows"]]
        self.assertEqual(parties, ["Amy's Bird Sanctuary", "Freeman Sporting Goods:0969 Ocean View Road"])

    def test_tally_stock_summary_asks_for_the_warehouse_then_opens(self):
        rows, _ = SHIPPED["tally_prime.stock_summary"]
        result, file_doc = self._check("tally_prime.stock_summary", rows, {})
        self.assertEqual(result["needs"][0]["name"], "warehouse")
        again = self.registry.get_tool("start_import_session")._safe_execute(
            {"file_id": file_doc.name, "values": {"warehouse": "Stores - X"}}
        )["result"]
        session = again["sessions"][0]
        self.assertEqual(session["engine"], "opening_stock")
        self.assertEqual(session["sample_rows"][0][:5], ["Steel Rod 10mm", "10", "Nos", "100", "1000"])
        self.assertEqual(session["columns"]["Warehouse"], "warehouse")

    def test_a_decimal_comma_csv_keeps_its_values_and_its_clean_copy_is_not_typed(self):
        # A Tally stock summary saved as CSV on a site that writes 1.500,50: comma decimals, quoted
        buffer = io.StringIO()
        csv.writer(buffer, lineterminator="\n").writerows(
            [
                ["Acme Pvt Ltd"],
                ["Stock Summary"],
                ["1-Apr-24 to 31-Mar-25"],
                ["Particulars", "Quantity", "Rate", "Value"],
                ["Raw Material", "", "", "3.188,56"],
                ["  Steel Rod 10mm", "2,125 Kg", "1.500,50/Kg", "3.188,56"],
                ["Grand Total", "", "", "3.188,56"],
            ]
        )
        file_doc = self._upload(f"stock-{self.tag}.csv", buffer.getvalue().encode("utf-8"))
        with patch.object(opening_stock, "decimal_comma", return_value=True):
            reply = self.registry.get_tool("start_import_session")._safe_execute(
                {"file_id": file_doc.name, "values": {"warehouse": "Stores - X"}}
            )
            self.assertTrue(reply["success"], reply)
            self.assertIn("sessions", reply["result"], reply)
            session = reply["result"]["sessions"][0]
            self.assertEqual(
                session["sample_rows"][0][:5], ["Steel Rod 10mm", "2.125", "Kg", "1500.5", "3188.56"]
            )
            # The clean copy holds numbers as Python writes them: its '2.125' is not the user's 2.125 thousand
            doc = frappe.get_doc("FAC Import Session", session["session_id"])
            clean = frappe.get_doc("File", doc.source_file)
            self.assertFalse(opening_stock.typed_file(doc, clean))
            cells = opening_stock.cell_texts(
                clean, ["2.125"], {0: "qty"}, opening_stock.typed_file(doc, clean)
            )
            self.assertEqual(opening_stock.number(cells["qty"]), 2.125)
            # A corrected CSV the user links later is their own typing again
            self.assertTrue(opening_stock.typed_file(doc, file_doc))

    def test_quickbooks_aging_columns_are_accepted_by_set_opening_invoices(self):
        if not opening_invoices.in_fiscal_year(frappe.utils.nowdate(), self._company()):
            self.skipTest("today is in no fiscal year of the test company")
        result, _ = self._check(
            "quickbooks_online.ar_aging_summary", *SHIPPED["quickbooks_online.ar_aging_summary"]
        )
        session = result["sessions"][0]
        reply = self.registry.get_tool("set_opening_invoices")._safe_execute(
            {
                "session": session["session_id"],
                "company": self._company(),
                "invoice_type": session["invoice_type"],
                "cut_off": frappe.utils.nowdate(),
                "columns": session["columns"],
            }
        )
        self.assertTrue(reply["success"], reply)

    def _company(self):
        return frappe.get_all("Company", pluck="name", order_by="creation", limit=1)[0]

    def test_zoho_contacts_ties_are_reported_for_the_model_to_ask(self):
        # Both Zoho contact presets have 2 required and 5 optional groups, so headers common to both tie.
        rows = [["Display Name", "EmailID", "Status"], ["Ravi Traders", "", "Active"]]
        file_doc = as_xlsx_file(self, f"zoho-contacts-{self.tag}.xlsx", rows)
        result = self.registry.get_tool("start_import_session")._safe_execute({"file_id": file_doc.name})[
            "result"
        ]
        self.assertEqual(result["preset"]["id"], "zoho_books.customers")
        self.assertIn("zoho_books.vendors", result["also_matched"])

    def test_a_named_preset_settles_the_zoho_contacts_tie(self):
        rows = [["Display Name", "EmailID", "Status"], ["Ravi Traders", "", "Active"]]
        file_doc = as_xlsx_file(self, f"zoho-contacts-{self.tag}.xlsx", rows)
        reply = self.registry.get_tool("start_import_session")._safe_execute(
            {"file_id": file_doc.name, "preset": "zoho_books.vendors"}
        )
        self.assertTrue(reply["success"], reply)
        result = reply["result"]
        self.assertEqual(result["preset"]["id"], "zoho_books.vendors")
        self.assertEqual(result["also_matched"], [])
        self.assertEqual([s["doctype"] for s in result["sessions"]], ["Supplier"])
        self.assertEqual(
            frappe.db.get_value("FAC Import Session", result["sessions"][0]["session_id"], "target_doctype"),
            "Supplier",
        )

    def _cleaned_rows(self, preset_id, rows):
        presets.clear_cache()
        found = next(
            m for m in presets.match_sheet(rows, presets.load_presets()) if m.preset["id"] == preset_id
        )
        target = presets.clean(presets.Sheet(None, rows), found).targets[0]
        return [dict(zip(target.header, row)) for row in target.rows]

    def test_quickbooks_desktop_accounts_are_flagged_as_groups_and_typed(self):
        rows = [
            ["Account", "Type", "Balance Total", "Description", "Accnt. #", "Tax Line"],
            ["1500 · Truck", "Fixed Asset", "0.00", "", "1500", ""],
            ["1500 · Truck:1510 · Depreciation", "Fixed Asset", "-2,500.00", "", "1510", ""],
            ["2100 · Visa", "Credit Card", "0.00", "", "2100", ""],
            ["2200 · Loan", "Long Term Liability", "0.00", "", "2200", ""],
            ["1800 · Deposits", "Other Asset", "0.00", "", "1800", ""],
            ["1200 · Prepaid", "Other Current Asset", "0.00", "", "1200", ""],
        ]
        by_name = {r["Account"]: r for r in self._cleaned_rows("quickbooks_desktop.account_listing", rows)}
        self.assertEqual(by_name["Truck"]["Is Group"], "1")
        self.assertEqual(by_name["Depreciation"]["Is Group"], "0")
        # Credit Card is not Payable: the opening journal leaves Payable to the opening invoices.
        self.assertEqual(by_name["Visa"]["Account Type"], "")
        self.assertEqual(by_name["Loan"]["Account Type"], "Liability")
        self.assertEqual(by_name["Deposits"]["Account Type"], "")
        self.assertEqual(by_name["Prepaid"]["Account Type"], "Current Asset")

    def test_zoho_chart_of_accounts_are_flagged_as_groups_and_typed(self):
        rows = [
            ["Account Name", "Account Type", "Account Code", "Parent Account"],
            ["Rent", "Expense", "6000", ""],
            ["Office Rent", "Expense", "6001", "Rent"],
            ["Other Receivables", "Other Current Asset", "1200", ""],
            ["Deferred Tax", "Deferred Tax Asset", "1300", ""],
            ["Corporate Card", "Credit Card", "2100", ""],
            ["Term Loan", "Long Term Liability", "2200", ""],
            ["Security Deposit", "Other Asset", "1400", ""],
        ]
        by_name = {r["Account"]: r for r in self._cleaned_rows("zoho_books.chart_of_accounts", rows)}
        self.assertEqual(by_name["Rent"]["Is Group"], "1")
        self.assertEqual(by_name["Office Rent"]["Is Group"], "0")
        self.assertEqual(by_name["Other Receivables"]["Account Type"], "Current Asset")
        self.assertEqual(by_name["Deferred Tax"]["Account Type"], "")
        self.assertEqual(by_name["Corporate Card"]["Account Type"], "")
        self.assertEqual(by_name["Term Loan"]["Account Type"], "Liability")
        self.assertEqual(by_name["Security Deposit"]["Account Type"], "")

    def test_a_busy_trial_balance_titled_as_on_is_still_busy(self):
        rows = [
            ["Acme Pvt Ltd"],
            ["As on 31-03-2025"],
            ["Account Name", "Debit", "Credit"],
            ["Cash", "200.00", ""],
            ["Capital", "", "200.00"],
            ["Total", "200.00", "200.00"],
        ]
        self._check("busy.trial_balance", rows, {"Trial balance": 2})

    def test_a_plain_file_is_not_claimed_by_any_preset(self):
        for content in (
            b"Account,Debit,Credit,Balance\nCash,100,,\n",
            b"Customer Name,Customer Group,Territory\nA,B,C\n",
            # Generic headers that untitled presets accept as required: not claimed without an
            # optional column of that system (signature.min_optional).
            b"Name,Group,Phone\nA,B,C\n",
            b"Name,Parent\nA,B\n",
            b"Name,Under\nA,B\n",
            b"Name,UOM\nA,Nos\n",
            b"Item Name,SKU\nA,B\n",
        ):
            with self.subTest(content.split(b"\n")[0]):
                file_doc = self._upload(f"plain-{self.tag}.csv", content)
                reply = self.registry.get_tool("start_import_session")._safe_execute(
                    {"file_id": file_doc.name}
                )
                self.assertTrue(reply["success"], reply)
                self.assertNotIn("matched", reply["result"])
                self.assertTrue(reply["result"]["session_id"], reply["result"])

    def test_tally_ledgers_customers_pass_the_dry_run(self):
        result, _ = self._check("tally_prime.ledgers", *SHIPPED["tally_prime.ledgers"])
        # Failures inside groups (a customer that already exists, say) are fine: the tool must not fail.
        customers = next(s for s in result["sessions"] if s["target"] == "Customers")
        reply = self.registry.get_tool("dry_run_import")._safe_execute({"session": customers["session_id"]})
        self.assertTrue(reply["success"], reply)
        self.assertEqual(reply["result"]["total_rows"], 1, reply)

    def test_quickbooks_aging_passes_set_opening_invoices_and_the_dry_run(self):
        if not opening_invoices.in_fiscal_year(frappe.utils.nowdate(), self._company()):
            self.skipTest("today is in no fiscal year of the test company")
        result, _ = self._check(
            "quickbooks_online.ar_aging_summary", *SHIPPED["quickbooks_online.ar_aging_summary"]
        )
        session = result["sessions"][0]
        setup = self.registry.get_tool("set_opening_invoices")._safe_execute(
            {
                "session": session["session_id"],
                "company": self._company(),
                "invoice_type": session["invoice_type"],
                "cut_off": frappe.utils.nowdate(),
                "columns": session["columns"],
            }
        )
        self.assertTrue(setup["success"], setup)
        reply = self.registry.get_tool("dry_run_import")._safe_execute({"session": session["session_id"]})
        self.assertTrue(reply["success"], reply)


class TestPresetIntoTheEngines(OpeningBalancesCase):
    """A cleaned trial balance goes through set_opening_balances and the dry run unchanged."""

    def test_quickbooks_trial_balance_balances_in_the_engine(self):
        rows = [
            ["Trial Balance"],
            ["Test Co"],
            ["As of December 31, 2025"],
            [],
            ["", "Debit", "Credit"],
            [self.bank, "1,201.00", ""],
            [self.capital, "", "1,201.00"],
            ["TOTAL", "$1,201.00", "$1,201.00"],
        ]
        file_doc = frappe.get_doc(
            {
                "doctype": "File",
                "file_name": f"qbo-tb-{self.tag}.xlsx",
                "content": xlsx(rows),
                "is_private": 1,
            }
        ).insert()
        presets.clear_cache()
        result = self._run("start_import_session", {"file_id": file_doc.name})["result"]
        session = result["sessions"][0]
        self.assertEqual(result["preset"]["id"], "quickbooks_online.trial_balance")
        self.assertEqual(session["rows"], 2)
        setup = self._setup(session["session_id"], columns=session["columns"])
        self.assertTrue(setup["success"], setup)
        checked = self._run("dry_run_import", {"session": session["session_id"]})
        self.assertTrue(checked["success"], checked)
        self.assertEqual(checked["result"]["groups"], [], checked)
