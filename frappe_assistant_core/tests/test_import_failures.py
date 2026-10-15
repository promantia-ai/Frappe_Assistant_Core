# Frappe Assistant Core - AI Assistant integration for Frappe Framework
# Copyright (C) 2025 Paul Clinton
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

"""
Fix failed rows and re-import only those (#33994): an import's failures are grouped and
explained as the dry run's are; once the cause is fixed, only the rows not imported are checked
and imported, in a new Data Import on the same session; the rest come as a file with the reason
for each; and that file, corrected, goes through the dry run before it is imported.

The failure is ERPNext's own: gamma's Customer Group is made a group node after the check (as
anyone in Desk can), and ERPNext refuses a customer in a group node. Nothing is mocked.
"""

import csv
import io
import unittest
from unittest.mock import patch

import frappe
import openpyxl

from frappe_assistant_core.plugins.core import import_fixes, import_run
from frappe_assistant_core.plugins.core.import_failures import REASON_COLUMN, ROW_COLUMN
from frappe_assistant_core.plugins.core.import_session import FileProblem, read_sheet_rows
from frappe_assistant_core.tests.test_import_run import MAPPING, ImportRunCase

GROUP_NODE = "Cannot select a Group type Customer Group"


@unittest.skipUnless("erpnext" in frappe.get_installed_apps(), "needs ERPNext's Customer")
class TestImportFailures(ImportRunCase):
    def setUp(self):
        super().setUp()
        self.held = f"Held {self.tag}"
        frappe.get_doc(
            {
                "doctype": "Customer Group",
                "customer_group_name": self.held,
                "parent_customer_group": "All Customer Groups",
            }
        ).insert()

    # -- helpers -----------------------------------------------------------------------

    def _partly_imported(self, file_doc=None, sheet=None):
        """Alpha and delta imported, gamma refused by ERPNext, beta left out at the check."""
        session, checked = self._ready_session(file_doc, sheet, map_to=self.held)
        # Production reaches this when someone changes the group in Desk between check and import.
        frappe.db.set_value("Customer Group", self.held, "is_group", 1)
        with self._each_row_commits():
            response = self._import(session, checked)
        self.assertTrue(response["success"], response)
        self.assertEqual(response["result"]["failed"], 1, response)
        return session, checked, response["result"]

    def _ungroup(self):
        fixed = self._run(
            "update_document", {"doctype": "Customer Group", "name": self.held, "data": {"is_group": 0}}
        )
        self.assertTrue(fixed["success"], fixed)

    def _check(self, session, **arguments):
        response = self._run("dry_run_import", {"session": session, **arguments})
        self.assertTrue(response["success"], response)
        return response["result"]

    def _failed_rows(self, session):
        out = self._check(session, failed_rows_file=True)
        file_doc = frappe.get_doc("File", {"file_url": out["file_url"]})
        content = file_doc.get_content()
        text = content.decode("utf-8") if isinstance(content, bytes) else content
        return out, file_doc, list(csv.reader(io.StringIO(text)))

    def _csv(self, name, rows):
        buffer = io.StringIO()
        csv.writer(buffer, lineterminator="\n").writerows(rows)
        return self._file(name, buffer.getvalue().encode())

    def _count(self, name):
        return frappe.db.count("Customer", {"customer_name": name})

    # -- criterion 1: failures grouped and explained as in the dry run ----------------------------

    def test_rows_that_failed_are_grouped_and_explained_as_the_dry_run_explains(self):
        session, _checked, result = self._partly_imported()

        [group] = result["failed_rows"]
        self.assertEqual((group["error"], group["rows"], group["count"]), ("ValidationError", [4], 1))
        self.assertIn(GROUP_NODE, group["reason"])
        self.assertTrue(group["explanation"].startswith(f"1 row rejected by ERPNext: {GROUP_NODE}"), group)
        self.assertIn(f"row 4: {GROUP_NODE}", result["message"])
        self.assertIn("dry_run_import", result["next"])
        self.assertIn("failed_rows_file", result["next"])
        self.assertEqual(self._doc(session).status, "Partially Completed")

    # -- criterion 2: only the failed rows are imported again ----------------------------------

    def test_once_the_cause_is_fixed_only_the_failed_rows_are_imported(self):
        session, _checked, first = self._partly_imported()
        n = self._names()
        before = {
            name: frappe.db.get_value("Customer", {"customer_name": name}, "modified")
            for name in (n["alpha"], n["delta"])
        }
        self._ungroup()

        rechecked = self._check(session)
        self.assertEqual(
            (rechecked["ready"], rechecked["imported"], rechecked["excluded"], rechecked["failed"]),
            (1, 2, 1, 0),
            rechecked,
        )
        self.assertIn("2 rows already imported are skipped", rechecked["message"])
        with self._each_row_commits():
            response = self._import(session, rechecked)

        self.assertTrue(response["success"], response)
        result = response["result"]
        self.assertEqual((result["created"], result["failed"]), (1, 0))
        self.assertEqual(self._imported(), {n["alpha"], n["gamma"], n["delta"]})
        self.assertEqual([self._count(n[k]) for k in ("alpha", "gamma", "delta")], [1, 1, 1])
        after = {
            name: frappe.db.get_value("Customer", {"customer_name": name}, "modified") for name in before
        }
        self.assertEqual(after, before)  # the records already in are not touched
        self.assertEqual(
            frappe.db.get_value("Customer", {"customer_name": n["gamma"]}, "customer_group"), self.held
        )
        doc = self._doc(session)
        self.assertEqual(doc.status, "Completed")
        self.assertEqual([d.purpose for d in doc.data_imports], ["Import", "Retry"])
        self.assertNotEqual(result["data_import"], first["data_import"])
        # The final report covers every run.
        self.assertEqual(result["created_in_all"], 3)
        self.assertEqual([(r["created"], r["failed"]) for r in result["runs"]], [(2, 1), (1, 0)])
        self.assertIn("In all, 3 records imported over 2 Data Imports", result["message"])

    def test_the_retry_is_a_new_file_of_only_the_rows_left_and_the_first_log_stays(self):
        session, _checked, first = self._partly_imported()
        n = self._names()
        self._ungroup()
        rechecked = self._check(session)
        with self._each_row_commits():
            retry = self._import(session, rechecked)["result"]

        di = frappe.get_doc("Data Import", retry["data_import"])
        file_doc = frappe.get_doc(
            "File", frappe.get_all("File", {"file_url": di.import_file}, pluck="name")[0]
        )
        rows = read_sheet_rows(file_doc)
        # Rows 2, 3 and 5 blank, so row 4 is still row 4; the map made since is in it.
        self.assertEqual([row for row in rows[1:] if any(row)], [[n["gamma"], self.held, "Company"]])
        self.assertEqual(rows[3], [n["gamma"], self.held, "Company"])
        logged = frappe.get_all("Data Import Log", {"data_import": di.name}, pluck="row_indexes")
        self.assertEqual(logged, ["[4]"])
        # Frappe's own Retry would delete the first run's failure: a new Data Import keeps it.
        self.assertEqual(
            frappe.db.count("Data Import Log", {"data_import": first["data_import"], "success": 0}), 1
        )
        self.assertEqual(self._doc(session).data_imports[-1].source_file, self._doc(session).source_file)

    def test_rows_imported_from_desk_since_the_check_make_it_stale(self):
        from frappe.core.doctype.data_import.data_import import start_import

        session, _checked, first = self._partly_imported()
        self._ungroup()
        rechecked = self._check(session)
        with self._each_row_commits():
            start_import(first["data_import"])  # Desk's own Retry of the first Data Import
        self.assertEqual(self._count(self._names()["gamma"]), 1)

        response = self._import(session, rechecked)  # the card still says 1 row to import

        self.assertFalse(response["success"], response)
        self.assertIn("changed since its last check", response["result"]["error"])
        self.assertEqual(self._count(self._names()["gamma"]), 1)

    def test_trying_again_without_checking_first_is_refused(self):
        session, checked, _first = self._partly_imported()
        self._ungroup()

        response = self._import(session, checked)

        self.assertFalse(response["success"], response)
        self.assertIn("dry_run_import", response["result"]["error"])
        self.assertEqual(len(self._doc(session).data_imports), 1)
        self.assertEqual(self._count(self._names()["gamma"]), 0)

    def test_when_the_rest_is_left_out_nothing_is_imported_again(self):
        session, _checked, _first = self._partly_imported()
        n = self._names()
        rechecked = self._check(session)  # gamma's group is still a group node
        self.assertEqual((rechecked["failed"], rechecked["imported"]), (1, 2), rechecked)
        [group] = rechecked["groups"]

        left = self._check(
            session, fixes=[{"group": group["id"], "action": "leave_out", "reason": "Customer on hold"}]
        )
        self.assertEqual((left["status"], left["ready"]), ("Completed", 0), left)
        self.assertIn("Nothing is left to import: 2 rows already imported, 2 rows left out.", left["message"])

        response = self._import(session, left)

        self.assertFalse(response["success"], response)
        self.assertIn("already imported", response["result"]["error"])
        self.assertEqual(len(self._doc(session).data_imports), 1)
        self.assertEqual([self._count(n[k]) for k in ("alpha", "delta", "gamma")], [1, 1, 0])

    # -- criterion 3: the rows not imported, as a file with the reasons ----------------------------

    def test_the_rows_not_imported_come_as_a_file_with_the_reason_for_each(self):
        session, _checked, _first = self._partly_imported()
        n = self._names()

        out, file_doc, rows = self._failed_rows(session)

        self.assertEqual(out["rows"], 2)
        self.assertTrue(out["download_url"].endswith(out["file_url"]))
        self.assertEqual(
            (file_doc.is_private, file_doc.attached_to_doctype, file_doc.attached_to_name),
            (1, "FAC Import Session", session),
        )
        self.assertEqual(rows[0], [*MAPPING, ROW_COLUMN, REASON_COLUMN])
        self.assertEqual(rows[1], [n["beta"], f"Nowhere {self.tag}", "Company", "3", "Left out: Closed shop"])
        self.assertEqual(rows[2][:4], [n["gamma"], self.held, "Company", "4"])  # with the map made in chat
        self.assertIn(GROUP_NODE, rows[2][4])
        self.assertEqual(len(rows), 3)

    def test_asking_again_replaces_the_file(self):
        session, _checked, _first = self._partly_imported()

        self._failed_rows(session)
        out, _file_doc, _rows = self._failed_rows(session)

        attached = frappe.get_all(
            "File",
            {"attached_to_name": session, "file_name": f"{session}-failed-rows.csv"},
            pluck="file_url",
        )
        self.assertEqual(attached, [out["file_url"]])

    def test_before_any_import_the_file_has_the_rows_the_check_failed(self):
        n = self._names()
        lines = [
            ",".join(MAPPING),
            f"{n['alpha']},{self.group},Company",
            f"{n['beta']},Nowhere {self.tag},Company",
        ]
        file_doc = self._file("customers.csv", ("\n".join(lines) + "\n").encode())
        session = self._run("start_import_session", {"file_url": file_doc.file_url})["result"]["session_id"]
        self._run("set_column_mapping", {"session": session, "mapping": MAPPING, "doctype": "Customer"})
        self.assertEqual(self._check(session)["status"], "Dry Run Failed")

        _out, _file_doc, rows = self._failed_rows(session)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:4], [n["beta"], f"Nowhere {self.tag}", "Company", "3"])
        self.assertEqual(rows[1][4], f"Customer Group 'Nowhere {self.tag}' doesn't exist.")

    def test_rows_imported_from_desk_since_the_last_check_are_not_in_the_file(self):
        from frappe.core.doctype.data_import.data_import import start_import

        session, _checked, first = self._partly_imported()
        self.assertEqual(self._check(session)["status"], "Dry Run Failed")  # gamma's group still a node
        self._ungroup()
        with self._each_row_commits():
            start_import(first["data_import"])  # Desk's own Retry gets gamma in

        _out, _file_doc, rows = self._failed_rows(session)

        self.assertEqual([row[3] for row in rows[1:]], ["3"])  # beta, left out; not gamma

    def test_the_file_never_holds_rows_of_a_file_the_user_cannot_read(self):
        secret = self._file(
            "salaries.csv", b"Customer Name,Customer Group,Customer Type\nCEO pay,99,Company\n"
        )
        user = self.make_throwaway_user("clerk", roles=("Sales Manager", "Desk User"))
        frappe.set_user(user)
        own = self._file("mine.csv", f"{','.join(MAPPING)}\nOwn {self.tag},{self.group},Company\n".encode())
        session = self._run("start_import_session", {"file_url": own.file_url})["result"]["session_id"]
        # What a PUT on /api/resource can do: read-only fields are only read-only in the form.
        doc = self._doc(session)
        doc.update({"source_file": secret.name, "status": "Completed", "excluded_rows": "[2]"})
        doc.save()

        response = self._run("dry_run_import", {"session": session, "failed_rows_file": True})

        self.assertFalse(response["success"], response)
        self.assertNotIn("CEO pay", frappe.as_json(response))
        frappe.set_user("Administrator")
        self.assertFalse(
            frappe.get_all("File", {"attached_to_name": session, "file_name": ["like", "%failed-rows%"]})
        )

    def test_a_session_pointed_at_a_file_the_user_cannot_read_is_not_checked(self):
        secret = self._file("salaries.csv", f"{','.join(MAPPING)}\nCEO pay,{self.group},Company\n".encode())
        frappe.set_user(self.make_throwaway_user("clerk", roles=("Sales Master Manager", "Desk User")))
        own = self._file("mine.csv", f"{','.join(MAPPING)}\nOwn {self.tag},{self.group},Company\n".encode())
        session, _checked = self._ready_session(own)
        doc = self._doc(session)
        doc.source_file = secret.name  # over the REST API
        doc.save()

        check = self._run("dry_run_import", {"session": session})

        self.assertFalse(check["success"], check)
        self.assertIn("don't have access", check["result"]["error"])
        self.assertNotIn("CEO pay", frappe.as_json(check))

    def test_whether_a_check_is_current_is_never_read_from_a_file_the_user_cannot_read(self):
        secret = self._file("salaries.csv", f"{','.join(MAPPING)}\nCEO pay,{self.group},Company\n".encode())
        frappe.set_user(self.make_throwaway_user("clerk", roles=("Sales Master Manager", "Desk User")))
        own = self._file("mine.csv", f"{','.join(MAPPING)}\nOwn {self.tag},{self.group},Company\n".encode())
        session, _checked = self._ready_session(own)
        doc = self._doc(session)
        doc.source_file = secret.name  # over the REST API
        doc.save()

        with self.assertRaisesRegex(FileProblem, "don't have access"):
            import_fixes.checked_as_is(self._doc(session))

    def test_with_every_row_in_there_is_nothing_to_download(self):
        n = self._names()
        file_doc = self._file("one.csv", f"{','.join(MAPPING)}\n{n['alpha']},{self.group},Company\n".encode())
        session, checked = self._ready_session(file_doc)
        with self._each_row_commits():
            self._import(session, checked)

        response = self._run("dry_run_import", {"session": session, "failed_rows_file": True})

        self.assertFalse(response["success"], response)
        self.assertIn("nothing to download", response["result"]["error"])

    def test_while_the_import_runs_there_is_no_file_yet(self):
        session, checked = self._ready_session()
        with patch.object(import_run, "enqueue", autospec=True):  # queued: the session is Importing
            self._import(session, checked)

        response = self._run("dry_run_import", {"session": session, "failed_rows_file": True})

        self.assertFalse(response["success"], response)
        self.assertIn("still running", response["result"]["error"])

    def test_the_file_is_asked_for_on_its_own(self):
        session, _checked, _first = self._partly_imported()

        response = self._run(
            "dry_run_import",
            {
                "session": session,
                "failed_rows_file": True,
                "fixes": [{"group": "G1", "action": "map", "to": "x"}],
            },
        )

        self.assertFalse(response["success"], response)
        self.assertIn("on its own", response["result"]["error"])

    # -- criterion 4: a corrected file goes through the dry run first ------------------------------

    def test_a_corrected_file_is_checked_and_then_imported_in_the_same_session(self):
        session, checked, first = self._partly_imported()
        n = self._names()
        _out, _file_doc, rows = self._failed_rows(session)
        rows[1][1] = self.group  # beta: a group that exists
        rows[2][1] = self.group  # gamma: one that isn't a group node
        corrected = self._csv("corrected.csv", rows)

        # Not imported before it is checked.
        self.assertFalse(self._import(session, checked)["success"])
        result = self._check(session, file_id=corrected.name)  # as uploaded in File Manager

        self.assertEqual(
            (result["status"], result["ready"], result["imported"]), ("Ready to Import", 2, 0), result
        )
        doc = self._doc(session)
        self.assertEqual(doc.source_file, corrected.name)
        mapping = frappe.parse_json(doc.template_options)["column_to_field_map"]
        self.assertEqual((mapping["3"], mapping["4"]), ("Don't Import", "Don't Import"))  # the added columns

        with self._each_row_commits():
            response = self._import(session, result)

        self.assertTrue(response["success"], response)
        self.assertEqual(self._imported(), set(n.values()))
        self.assertEqual([self._count(name) for name in n.values()], [1, 1, 1, 1])
        doc = self._doc(session)
        self.assertEqual([d.purpose for d in doc.data_imports], ["Import", "Import"])
        self.assertEqual(
            [d.source_file for d in doc.data_imports],
            [first_file := doc.data_imports[0].source_file, corrected.name],
        )
        self.assertNotEqual(first_file, corrected.name)
        self.assertEqual(response["result"]["created_in_all"], 4)
        self.assertEqual(doc.status, "Completed")

    def test_rows_left_out_of_a_completed_import_come_back_corrected(self):
        session, checked = self._ready_session()  # beta left out, the rest imported
        n = self._names()
        with self._each_row_commits():
            self._import(session, checked)
        self.assertEqual(self._doc(session).status, "Completed")
        _out, _file_doc, rows = self._failed_rows(session)
        self.assertEqual([row[3:] for row in rows[1:]], [["3", "Left out: Closed shop"]])
        rows[1][1] = self.group
        corrected = self._csv("corrected.csv", rows)

        result = self._check(session, file_url=corrected.file_url)

        self.assertEqual((result["status"], result["ready"]), ("Ready to Import", 1), result)
        with self._each_row_commits():
            response = self._import(session, result)
        self.assertTrue(response["success"], response)
        self.assertEqual(self._imported(), set(n.values()))
        self.assertEqual(response["result"]["created_in_all"], 4)

    def test_checking_a_completed_session_again_leaves_it_completed(self):
        session, checked = self._ready_session()
        with self._each_row_commits():
            self._import(session, checked)

        result = self._check(session)

        self.assertEqual((result["status"], result["ready"], result["imported"]), ("Completed", 0, 3), result)
        self.assertIn("Nothing is left to import", result["message"])
        self.assertFalse(self._import(session, result)["success"])

    def test_a_file_sent_while_the_session_cannot_be_checked_is_not_kept(self):
        session, checked = self._ready_session()
        with patch.object(import_run, "enqueue", autospec=True):  # queued: the session is Importing
            self._import(session, checked)
        source = self._doc(session).source_file
        other = self._csv("other.csv", [list(MAPPING), ["Zed", self.group, "Company"]])

        response = self._run("dry_run_import", {"session": session, "file_url": other.file_url})

        self.assertFalse(response["success"], response)
        self.assertEqual(self._doc(session).source_file, source)

    def test_a_session_imported_before_runs_kept_their_file_takes_a_corrected_one(self):
        session, _checked, _first = self._partly_imported()
        # Sessions imported before source_file was kept have it empty on their Data Import rows.
        link = self._doc(session).data_imports[0]
        frappe.db.set_value(link.doctype, link.name, "source_file", None)
        _out, _file_doc, rows = self._failed_rows(session)
        rows[1][1] = rows[2][1] = self.group
        corrected = self._csv("corrected.csv", rows)

        result = self._check(session, file_url=corrected.file_url)

        # Rows 2 and 3 of the corrected file aren't rows 2 and 5 the first file imported.
        self.assertEqual((result["ready"], result["imported"]), (2, 0), result)

    def test_a_corrected_file_that_brings_back_an_imported_row_is_stopped_by_the_check(self):
        session, _checked, _first = self._partly_imported()
        n = self._names()
        _out, _file_doc, rows = self._failed_rows(session)
        rows[1][1] = rows[2][1] = self.group
        rows.append([n["alpha"], self.group, "Company", "2", ""])  # already imported
        corrected = self._csv("corrected.csv", rows)

        result = self._check(session, file_url=corrected.file_url)

        self.assertEqual((result["status"], result["failed"]), ("Dry Run Failed", 1), result)
        self.assertEqual(result["groups"][0]["cause"], "Already in ERPNext")

    def test_a_corrected_file_must_keep_the_columns(self):
        session, _checked, _first = self._partly_imported()
        other = self._csv("other.csv", [["Name", "Group"], ["Zed", self.group]])

        response = self._run("dry_run_import", {"session": session, "file_url": other.file_url})

        self.assertFalse(response["success"], response)
        self.assertIn("columns don't match", response["result"]["error"])
        self.assertNotEqual(self._doc(session).source_file, other.name)

    def _workbook_partly_imported(self):
        """The same, from the "Customers" sheet of a workbook: alpha in, gamma refused."""
        n = self._names()
        book = openpyxl.Workbook()
        book.active.title = "Notes"
        book.active.append(["Anything"])
        sheet = book.create_sheet("Customers")
        sheet.append(list(MAPPING))
        sheet.append([n["alpha"], self.group, "Company"])
        sheet.append([n["gamma"], f"Retial {self.tag}", "Company"])
        buffer = io.BytesIO()
        book.save(buffer)
        return self._partly_imported(self._file("customers.xlsx", buffer.getvalue()), "Customers")

    def test_a_workbook_session_takes_its_corrected_rows_back_as_csv(self):
        n = self._names()
        session, _checked, _first = self._workbook_partly_imported()
        _out, _file_doc, rows = self._failed_rows(session)
        rows[1][1] = self.group
        corrected = self._csv("corrected.csv", rows)

        result = self._check(session, file_url=corrected.file_url)

        self.assertEqual((result["status"], result["ready"]), ("Ready to Import", 1), result)
        self.assertIsNone(self._doc(session).sheet)  # a CSV has no sheets
        with self._each_row_commits():
            response = self._import(session, result)
        self.assertTrue(response["success"], response)
        self.assertEqual(self._imported(), {n["alpha"], n["gamma"]})

    def test_a_corrected_file_saved_as_a_workbook_is_read_whatever_its_sheet_is_called(self):
        session, _checked, _first = self._workbook_partly_imported()
        _out, _file_doc, rows = self._failed_rows(session)
        rows[1][1] = self.group
        book = openpyxl.Workbook()
        book.active.title = f"{session}-failed-rows"  # Excel's name for it, opened from the CSV
        for row in rows:
            book.active.append(row)
        buffer = io.BytesIO()
        book.save(buffer)

        result = self._check(session, file_url=self._file("corrected.xlsx", buffer.getvalue()).file_url)

        self.assertEqual((result["status"], result["ready"]), ("Ready to Import", 1), result)
        self.assertEqual(self._doc(session).sheet, f"{session}-failed-rows")

    # -- Frappe v16 refuses SQL functions in get_all's fields ------------------------------------

    def test_the_runs_are_counted_without_asking_get_all_for_a_sql_function(self):
        from frappe_assistant_core.plugins.core import import_failures
        from frappe_assistant_core.tests.test_import_opening_stock import refuse_sql_functions

        session, _checked, _result = self._partly_imported()

        with patch("frappe.get_all", refuse_sql_functions(frappe.get_all)):
            runs = import_failures.runs(self._doc(session))

        self.assertEqual([(run["created"], run["failed"]) for run in runs], [(2, 1)])
