# Frappe Assistant Core - AI Assistant integration for Frappe Framework
# Copyright (C) 2025 Paul Clinton
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

"""
Helpers for start_import_session: checking file access, checking that the user
may import into a target DocType, and reading a spreadsheet.

Cells are read as text, the way Data Import will see them. The assistant gets the
headers, row counts and a few sample rows; the session stores only the headers and
row counts. The sheet it works on is read once into a private CSV the session owns, and
every later step reads that copy (#34405).
"""

import codecs
import csv
import datetime
import inspect
import io
import itertools
import os
import re
import time
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple
from urllib.parse import urlparse

import frappe
import requests
from frappe import _
from frappe.core.api.file import get_max_file_size
from frappe.utils import cint

from frappe_assistant_core.chat.api.settings.uploads import CHAT_UPLOAD_MAX_BYTES

SUPPORTED_EXTENSIONS = ("csv", "xlsx", "xls")
SAMPLE_SIZE = 10
MAX_FILE_SIZE = CHAT_UPLOAD_MAX_BYTES  # the chat's ceiling; the site's Max File Size may be lower
# A Google Sheet's download: seconds to connect and between two pieces, and in all. The whole is
# checked as a piece arrives, so it can run one piece's wait over: 60 + 15 stays well inside a web
# worker's 120 s timeout, which would otherwise end the request before the reply (#34438).
DOWNLOAD_TIMEOUT = (10, 15)
DOWNLOAD_SECONDS = 60
HEADER_SEARCH_ROWS = 20  # how far down a sheet its headings are looked for (#34403)
FIRST_ROWS_SHOWN = 5  # a sheet's top rows, shown when its headings can't be found for sure
CELL_CHARS_SHOWN = 200  # a cell shown to the model is cut to this many characters (#34406)
DELIMITERS = (",", ";", "\t", "|")  # a CSV's, in the order a tie goes


@dataclass
class SheetSummary:
    """What the assistant gets to see of one sheet."""

    name: Optional[str]  # None for a .csv file
    columns: List[str] = field(default_factory=list)
    row_count: int = 0
    sample_rows: List[List[str]] = field(default_factory=list)
    header_row: int = 1  # the sheet's row holding the headings, counted from 1 (#34403)
    header_sure: bool = True  # False when rows above the headings hold several values too
    first_rows: List[List[str]] = field(default_factory=list)  # the top rows, when not on row 1
    rows: Optional[List[List[str]]] = None  # the headings and every row under them, when kept

    def as_dict(self) -> Dict[str, Any]:
        """What the session stores: no sample rows, so no data from the file."""
        return {"name": self.name, "columns": self.columns, "row_count": self.row_count}


@dataclass
class SpreadsheetRead:
    sheets: List[SheetSummary]  # every sheet, in workbook order
    chosen: SheetSummary  # the sheet the session works on


class FileProblem(Exception):
    """A file that cannot be opened as a session. The message is shown to the user."""


# -- files -----------------------------------------------------------------------


def get_readable_file(file_url: str):
    """Return the File the current user may read at ``file_url``.

    Identical uploads share one ``file_url`` across several File records (Frappe
    stores the content once), so every record behind the URL is tried, and the
    user gets the first one they may read.
    """
    for name in frappe.get_all("File", filters={"file_url": file_url}, pluck="name"):
        file_doc = frappe.get_doc("File", name)
        if frappe.has_permission("File", "read", doc=file_doc):
            return file_doc
    # One message for "missing" and "not yours", so file names cannot be probed.
    raise FileProblem(_("File not found, or you don't have access to it."))


def get_readable_file_by_id(file_id: str):
    """Return the File named ``file_id`` (its ID in File Manager), if the current user may read it."""
    if frappe.db.exists("File", {"name": file_id, "is_folder": 0}):
        file_doc = frappe.get_doc("File", file_id)
        if frappe.has_permission("File", "read", doc=file_doc):
            return file_doc
    raise FileProblem(_("File not found, or you don't have access to it."))


def find_site_file(file_id: str = "", file_url: str = ""):
    """The File the current user may read, by its ID or its URL. Both may be given when they name
    the same file, as a chat attachment gives them."""
    if file_id:
        file_doc = get_readable_file_by_id(file_id)
        if file_url and file_doc.file_url != file_url:
            raise FileProblem(_("file_id and file_url name different files: give one of them."))
        return file_doc
    return get_readable_file(file_url)


def file_from_google_sheet(url: str):
    """A private .csv copy, owned by the user, of the Google Sheets tab ``url`` links to (#34006).

    The tab in the link is downloaded, else the first. The session works on the copy like on any
    file of this site, so what the dry run checks is what gets imported, whatever is changed in
    the sheet meanwhile.
    """
    from frappe.utils.csvutils import validate_google_sheets_url

    url = (url or "").strip()
    try:
        validate_google_sheets_url(url)
    except frappe.ValidationError:
        frappe.clear_messages()  # Frappe's pop-up says the same, less helpfully
        raise FileProblem(
            _(
                "This isn't a Google Sheets link. Open the sheet in Google Sheets and copy the link "
                "from the browser's address bar."
            )
        )
    try:
        content = _download_sheet(url)
    except (requests.ConnectionError, requests.Timeout):
        raise FileProblem(_cant_reach_google())
    except requests.HTTPError as e:
        if e.response is not None and e.response.status_code >= 500:
            raise FileProblem(_cant_reach_google())
        raise FileProblem(_cant_open_sheet())  # a deleted or mistyped sheet (404), a wrong tab (400)
    except requests.RequestException:
        raise FileProblem(_cant_open_sheet())

    sheet_id = re.sub(r"[^A-Za-z0-9_-]", "", urlparse(url).path.split("/d/", 1)[-1].split("/", 1)[0])
    return frappe.get_doc(
        {
            "doctype": "File",
            "file_name": f"google-sheet-{sheet_id[:8] or 'copy'}.csv",
            "content": content,
            "is_private": 1,
        }
    ).insert()


def _download_sheet(url: str) -> bytes:
    """The CSV export of the tab ``url`` links to, the way Frappe's
    get_csv_content_from_google_sheets builds its link, but streamed: refused once it passes
    max_file_size() or DOWNLOAD_SECONDS, so a huge or trickling sheet can't hold a worker."""
    # The tab's number only: a "link to this range" adds "&range=…" after it
    gid = re.search(r"gid=(\d+)", url)
    export = url.split("/edit", 1)[0] + f"/export?format=csv&gid={gid.group(1) if gid else 0}"
    deadline = time.monotonic() + DOWNLOAD_SECONDS
    limit = max_file_size()
    too_large = _("The sheet is larger than {0} MB, the largest file this site takes.").format(
        round(limit / (1024 * 1024), 1)
    )

    with requests.get(
        export, headers={"Accept": "text/csv"}, stream=True, timeout=DOWNLOAD_TIMEOUT
    ) as response:
        response.raise_for_status()
        if cint(response.headers.get("Content-Length")) > limit:
            raise FileProblem(too_large)
        content = bytearray()
        # As the pieces arrive: a large piece trickling in would put off the time limit
        for chunk in response.iter_content(chunk_size=None):
            content += chunk
            if len(content) > limit:
                raise FileProblem(too_large)
            if time.monotonic() > deadline:
                raise FileProblem(_cant_reach_google())

    # A sheet that isn't shared answers 200 with Google's sign-in page.
    if bytes(content).strip().endswith(b"</html>"):
        raise FileProblem(_cant_open_sheet())
    return bytes(content)


def discard_sheet_copy(file_doc) -> None:
    """Delete a Google Sheet's copy no session uses. The disk file stays while another File
    shares its content."""
    if file_doc:
        # made by this same call, for this user
        frappe.delete_doc("File", file_doc.name, ignore_permissions=True)


def _cant_open_sheet() -> str:
    return _(
        "I can't open this Google Sheet. Share it so anyone with the link can view it: in Google "
        "Sheets click Share, under General access choose Anyone with the link (Viewer), then paste "
        "the link again. If it is shared, open the tab to import and copy the link from the "
        "browser's address bar."
    )


def _cant_reach_google() -> str:
    return _("Google Sheets didn't answer. Please try again in a moment.")


def readable_session_file(session):
    """The session's file, when the current user may read it. A session's fields can be written
    over the REST API, read-only or not, so its source_file is checked each time it is read."""
    if session.source_file and frappe.db.exists("File", session.source_file):
        file_doc = frappe.get_doc("File", session.source_file)
        if frappe.has_permission("File", "read", doc=file_doc):
            return file_doc
    raise FileProblem(_("File not found, or you don't have access to it."))


def read_spreadsheet(
    file_doc, sheet: Optional[str] = None, header_row: Optional[int] = None, keep_rows: bool = False
) -> SpreadsheetRead:
    """Summarise every sheet of ``file_doc`` and choose the one to work on.

    The chosen sheet is ``sheet`` when given, else the first sheet with rows. Its headings are on
    ``header_row`` (counted from 1) when given, else found below any title row (#34403). With
    ``keep_rows``, the chosen sheet keeps its rows, for the session's own copy.
    """
    extension = (file_doc.get_extension()[1] or "").lstrip(".").lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise FileProblem(_("Only .csv, .xlsx or .xls files can be imported."))
    check_size(file_doc)

    try:
        content = file_bytes(file_doc)
        # Sheets stream past, so a large workbook is never held in memory as rows: only the
        # sheet the session may work on is kept.
        sheets: List[SheetSummary] = []
        for name, rows in _iter_sheets(content, extension):
            found = any(s.columns and s.row_count for s in sheets)
            mine = extension == "csv" or (name == sheet if sheet else not found)
            sheets.append(_summarise(name, rows, header_row if mine else None, keep_rows and mine))
        chosen = _choose(sheets, sheet if extension != "csv" else None)
    except FileProblem:
        raise
    except Exception as e:
        raise FileProblem(_("Could not read the file: {0}").format(_describe_read_error(e)))

    for summary in sheets:
        if summary is not chosen:
            summary.rows = None
    return SpreadsheetRead(sheets=sheets, chosen=chosen)


def max_file_size() -> int:
    """The largest file an import reads (#34413): the chat's ceiling, or System Settings' Max File
    Size when it is lower (Frappe's File refuses more; 25 MB by default)."""
    return min(MAX_FILE_SIZE, get_max_file_size())


def check_size(file_doc) -> None:
    """Raises FileProblem when ``file_doc`` is larger than max_file_size()."""
    limit = max_file_size()
    if (file_doc.file_size or 0) > limit:
        raise FileProblem(
            _("The file is larger than {0} MB, the largest file this site takes.").format(
                f"{limit / (1024 * 1024):g}"
            )
        )


def read_sheet_rows(file_doc, sheet: Optional[str] = None) -> List[List[str]]:
    """Every row of one sheet as text, header first, for the engine (never the model).

    Blank rows stay in place as empty rows, so a row's position is its line in the
    spreadsheet, the number Data Import reports and the user sees.
    """
    extension = (file_doc.get_extension()[1] or "").lstrip(".").lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise FileProblem(_("Only .csv, .xlsx or .xls files can be imported."))
    check_size(file_doc)
    try:
        content = file_bytes(file_doc)
        names = []
        for name, rows in _iter_sheets(content, extension):
            names.append(name)
            if extension == "csv" or not sheet or name == sheet:
                return _all_text_rows(rows)
    except FileProblem:
        raise
    except Exception as e:
        raise FileProblem(_("Could not read the file: {0}").format(_describe_read_error(e)))
    raise FileProblem(
        _("Sheet '{0}' not found. Sheets in this file: {1}.").format(
            sheet, ", ".join(str(name) for name in names)
        )
    )


def file_bytes(file_doc) -> bytes:
    """The file as it is stored. Frappe's get_content decodes text itself: v15 as UTF-8 only,
    v16 trying windows-1250 too, which reads UTF-16 as garbage (#34404)."""
    if "encodings" in inspect.signature(file_doc.get_content).parameters:
        content = file_doc.get_content(encodings=())  # Frappe 16: no decoding
    else:
        content = file_doc.get_content()  # Frappe 15: bytes, unless they are UTF-8
    return content.encode("utf-8") if isinstance(content, str) else content


def decode_csv(content: bytes, fallbacks: Tuple[str, ...] = ("windows-1250", "windows-1252")) -> str:
    """A CSV's text (#34404). A byte-order mark names UTF-16, as Excel's "Unicode Text" writes
    it, or UTF-8; else UTF-8, then each fallback in turn, the last of which never fails."""
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", *fallbacks[:-1]):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode(fallbacks[-1], errors="replace")


def csv_delimiter(text: str) -> str:
    """The CSV's delimiter: a comma, a semicolon (Excel in much of Europe), a tab or a pipe (#34404).

    The one that splits the most of the first rows into the same number of columns, more than
    one. csv.Sniffer gives up on a title row above the table, which has no delimiter at all.
    """
    sample = text[:65536].splitlines()[:50]
    best, best_rows = ",", 0
    for delimiter in DELIMITERS:
        counts = [len(row) for row in csv.reader(sample, delimiter=delimiter) if any(c.strip() for c in row)]
        if not counts:
            continue
        width = max(set(counts), key=lambda n: (counts.count(n), n))
        if width > 1 and counts.count(width) > best_rows:
            best, best_rows = delimiter, counts.count(width)
    return best


def save_session_copy(file_name: str, chosen: SheetSummary):
    """The sheet a session works on, as a private CSV of its own (#34405): its headings found and
    made unique, its text decoded, written once. Every later step reads this copy, so the file the
    user gave may go (chat retention deletes a chat's attachments) while the session carries on."""
    from frappe.core.doctype.file.exceptions import MaxFileSizeReachedError

    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(chosen.rows or [chosen.columns])
    content = buffer.getvalue().encode("utf-8")
    try:
        return frappe.get_doc(
            {
                "doctype": "File",
                # never the name of a file the session writes itself, which is replaced by name:
                # a corrected failed-rows file, uploaded again, is called the same as the next one
                "file_name": f"{os.path.splitext(file_name or '')[0] or 'import'} - session copy.csv",
                "content": content,
                "is_private": 1,
            }
        ).insert()
    except MaxFileSizeReachedError:
        # a workbook grows several times over as CSV: the copy, not the file, passed the limit
        frappe.clear_messages()
        raise FileProblem(
            _(
                "The sheet is {0} MB as text, more than the {1} MB this site takes for a file. Split it "
                "into smaller files, or ask an administrator to raise Max File Size in System Settings."
            ).format(f"{len(content) / (1024 * 1024):.1f}", f"{max_file_size() / (1024 * 1024):g}")
        ) from None


def attach_session_copy(copy, session_name: str) -> None:
    """Attach the session's own copy to it, so it lives and goes with the session."""
    frappe.db.set_value(
        "File", copy.name, {"attached_to_doctype": "FAC Import Session", "attached_to_name": session_name}
    )


def discard_session_copy(copy) -> None:
    """Delete a session copy no session took. The disk file stays while another File shares it."""
    if copy:
        frappe.delete_doc("File", copy.name, force=True, ignore_permissions=True)


def _all_text_rows(rows: Iterable) -> List[List[str]]:
    width = None
    out = []
    for raw in rows:
        values = [_as_text(v) for v in raw]
        if width is None:
            width = max((i + 1 for i, v in enumerate(values) if v), default=0)
        row = values[:width]
        out.append(row + [""] * (width - len(row)))
    return out


def _choose(sheets: List[SheetSummary], sheet: Optional[str]) -> SheetSummary:
    if sheet:
        for summary in sheets:
            if summary.name == sheet:
                if not summary.row_count:
                    raise FileProblem(_("Sheet '{0}' has no rows to import.").format(sheet))
                return summary
        names = ", ".join(s.name for s in sheets)
        raise FileProblem(_("Sheet '{0}' not found. Sheets in this file: {1}.").format(sheet, names))

    for summary in sheets:
        if summary.columns and summary.row_count:
            return summary
    raise FileProblem(_("The file has no rows to import."))


def _describe_read_error(e: Exception) -> str:
    import zipfile

    if isinstance(e, zipfile.BadZipFile):
        return _("it is not a valid Excel workbook")
    return str(e) or type(e).__name__


def _iter_sheets(content: bytes, extension: str) -> Iterator[Tuple[Optional[str], Iterable]]:
    """Yield (sheet name, rows) pairs. Rows hold raw cell values."""
    if extension == "csv":
        yield None, _csv_cells(content)
    elif extension == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                yield worksheet.title, worksheet.iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        for xls_sheet in book.sheets():
            yield xls_sheet.name, (_xls_row(book, xls_sheet, i) for i in range(xls_sheet.nrows))


def _csv_cells(content: bytes) -> Iterator[List[str]]:
    """A CSV's rows, decoded and split on its own delimiter, cells trimmed as Frappe trims them."""
    text = decode_csv(content)
    for row in csv.reader(io.StringIO(text, newline=""), delimiter=csv_delimiter(text)):
        yield [cell.strip() for cell in row]


def _xls_row(book, xls_sheet, index: int) -> list:
    import xlrd

    values = []
    for cell in xls_sheet.row(index):
        if cell.ctype == xlrd.XL_CELL_DATE:
            values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
            values.append(bool(cell.value))
        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            values.append(None)
        else:
            values.append(cell.value)
    return values


def _heading_index(top: List[List[str]]) -> Tuple[int, bool]:
    """Where a sheet's headings are among its top rows (0-based), and whether that is sure (#34403).

    Exports from Tally, Busy or Zoho put a title (the company, the report) above them. The
    headings are the first row filling more than half the table's width, taken both as most rows'
    width and as the widest row's. When the two disagree (a few "Period: ..." rows above a short
    table, say), or rows above the guess hold several values, it isn't sure: the user is asked.
    """
    filled = [sum(1 for cell in row if cell.strip()) for row in top]
    counts = [n for n in filled if n]
    if not counts:
        return 0, True
    most = max(set(counts), key=lambda n: (counts.count(n), n))
    by_most, by_widest = (next(i for i, n in enumerate(filled) if n * 2 > w) for w in (most, max(counts)))
    return by_widest, by_most == by_widest and all(n <= 1 for n in filled[:by_widest])


def _summarise(
    name: Optional[str], rows: Iterable, header_row: Optional[int] = None, keep: bool = False
) -> SheetSummary:
    """One sheet: its headings (on ``header_row``, else found), then every non-blank data row
    counted, as text cut to the headings' width. Trailing empty heading cells are formatting, not
    columns. With ``keep``, the headings and every row under them are kept, blank rows in place."""
    from frappe_assistant_core.utils.import_dry_run import unique_headers

    rows = iter(rows)
    top = [
        [_as_text(v) for v in raw] for raw in itertools.islice(rows, max(HEADER_SEARCH_ROWS, header_row or 0))
    ]
    if header_row:
        index, sure = header_row - 1, True
        if not any(any(cell.strip() for cell in row) for row in top):
            index = 0  # an empty sheet, before the one with rows
        elif index >= len(top) or not any(cell.strip() for cell in top[index]):
            raise FileProblem(
                _("Row {0} is blank: give the row that holds the column headings.").format(header_row)
            )
    else:
        index, sure = _heading_index(top)
    summary = SheetSummary(name=name, header_row=index + 1, header_sure=sure)
    if index or not sure:
        summary.first_rows = [
            [shown_cell(cell) for cell in _trimmed(row)] for row in top[: max(FIRST_ROWS_SHOWN, index + 2)]
        ]
    kept: Optional[List[List[str]]] = [] if keep else None

    width = None
    for values in itertools.chain(top[index:], ([_as_text(v) for v in raw] for raw in rows)):
        if width is None:
            width = max((i + 1 for i, v in enumerate(values) if v), default=0)
            summary.columns = [str(title) for title in unique_headers(values[:width])]
            if kept is not None:
                kept.append(summary.columns)
            continue
        row = values[:width]
        row += [""] * (width - len(row))
        if kept is not None:
            kept.append(row)
        if not any(values):
            continue
        summary.row_count += 1
        if len(summary.sample_rows) < SAMPLE_SIZE:
            summary.sample_rows.append([shown_cell(cell) for cell in row])

    if not summary.columns:
        summary.row_count = 0
        summary.sample_rows = []
    if kept is not None:
        while len(kept) > 1 and not any(kept[-1]):
            kept.pop()  # formatted but empty rows at the end of a sheet
        summary.rows = kept
    return summary


def shown_cell(value: str) -> str:
    """A cell as the model is shown it: cut to CELL_CHARS_SHOWN characters, so one long cell
    can't fill the reply (#34406). The file and its copy keep it whole."""
    return value if len(value) <= CELL_CHARS_SHOWN else value[:CELL_CHARS_SHOWN] + "…"


def _trimmed(row: List[str]) -> List[str]:
    """A row without its trailing empty cells."""
    end = max((i + 1 for i, v in enumerate(row) if v.strip()), default=0)
    return row[:end]


# A cell starting with one of these is a formula to a spreadsheet: the failed-rows file writes it
# with a leading ' (import_failures.csv_cell), which comes off again when it is read back (#34413).
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _as_text(value: Any) -> str:
    """A cell as Data Import would read it, written as plain text."""
    if value is None:
        return ""
    if (
        isinstance(value, str)
        and value.startswith("'")
        and value[1:2]
        and value[1:].startswith(FORMULA_START)
    ):
        return value[1:]
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# -- how many rows one import may have ----------------------------------------------

DEFAULT_MAX_IMPORT_ROWS = 50000


#: A session holds its file while it is being worked on (#34437): a check, an import, a retry of
#: failed rows (after a Failed import too). Once it is Completed, Undone or Cancelled, the file is
#: free again.
HOLDING_STATUSES = (
    "File Read",
    "Mapped",
    "Dry Run Failed",
    "Ready to Import",
    "Importing",
    "Partially Completed",
    "Failed",
    "Undoing",
    "Partially Undone",
)


def holding_session(file_name: str, user: Optional[str] = None) -> Optional[str]:
    """The open import session of ``user`` (default: the current one) that holds File
    ``file_name``, or None (#34003, #34437). The import tools read that file, showing the model its
    headers, samples and summaries only, never all its rows. Both are held: the file the user gave,
    and the session's own copy of it (#34405).

    Frappe keeps identical bytes once, so the same file uploaded again by the same user (a new File
    with the same file_url) is held too. Another user's upload of the same bytes is theirs to read,
    and their reply never names this user's session."""
    file_url = frappe.db.get_value("File", file_name, "file_url")
    copies = frappe.get_all("File", filters={"file_url": file_url}, pluck="name") if file_url else []
    names = list({file_name, *copies})
    for fieldname in ("source_file", "original_file"):
        held = frappe.db.get_value(
            "FAC Import Session",
            {
                fieldname: ["in", names],
                "user": user or frappe.session.user,
                "status": ["in", HOLDING_STATUSES],
            },
            "name",
        )
        if held:
            return held
    return None


def max_import_rows() -> int:
    """Assistant Core Settings' Max Rows per Import; 0 or blank is the default, never unlimited."""
    try:
        value = frappe.db.get_single_value("Assistant Core Settings", "max_import_rows")
    except frappe.ValidationError:
        # the code is newer than the site: `bench migrate` hasn't added the field yet
        frappe.clear_messages()
        value = None
    return cint(value) or DEFAULT_MAX_IMPORT_ROWS


def row_limit_problem(row_count: int) -> Optional[str]:
    """Why a file of ``row_count`` rows can't be imported in one go, or None (#34003)."""
    limit = max_import_rows()
    if cint(row_count) <= limit:
        return None
    return _(
        "The file has {0} rows, more than the {1} one import may have. Split it into files of at "
        "most {1} rows and import them one after another. (The limit is Max Rows per Import in "
        "Assistant Core Settings.)"
    ).format(row_count, limit)


# -- target DocType ----------------------------------------------------------------


def import_target_problem(doctype: str, import_type: str = "Insert New Records") -> Optional[str]:
    """Why the current user cannot import into ``doctype``, or None when they can.

    Data Import itself requires the DocType to allow import and the user to hold
    Import permission on it; inserting each row also needs Create, updating needs Write.
    """
    if not frappe.db.exists("DocType", doctype):
        return _("DocType {0} does not exist.").format(doctype)

    meta = frappe.get_meta(doctype)
    if meta.istable:
        return _("{0} is a child table. Import it through its parent DocType.").format(doctype)
    if meta.issingle:
        return _("{0} is a single settings record, not a list, so it can't be imported.").format(doctype)
    if not meta.allow_import:
        return _("{0} can't be imported: Allow Import is turned off for it.").format(doctype)
    if import_type == "Update Existing Records":
        if not frappe.has_permission(doctype, "write"):
            return _("You don't have permission to change {0} records, so you can't update them.").format(
                doctype
            )
    elif not frappe.has_permission(doctype, "create"):
        return _("You don't have permission to create {0} records, so you can't import them.").format(doctype)
    if not frappe.has_permission(doctype, "import"):
        return _(
            "You can create {0} records but don't have Import permission on {0}. "
            "A System Manager can grant it in Role Permissions Manager."
        ).format(doctype)
    return None
