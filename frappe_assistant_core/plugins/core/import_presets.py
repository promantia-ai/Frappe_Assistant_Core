# Frappe Assistant Core - AI Assistant integration for Frappe Framework
# Copyright (C) 2025 Paul Clinton
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.

"""
Import presets (#34000): recognise a source system's export by its headers, and clean it.

A preset is data: the headers that identify an export, how each column is read, and which
ERPNext record or opening import the rows become. FAC ships presets in data/import_presets/;
other apps add theirs through the assistant_import_presets hook. This module reads, matches and
cleans. It writes nothing: apply_import_preset writes the cleaned files and opens the sessions.
"""

import csv
import datetime
import fnmatch
import glob
import io
import itertools
import json
import os
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple

import frappe
from frappe import _
from openpyxl import load_workbook

from frappe_assistant_core.plugins.core.import_session import (
    MAX_FILE_SIZE,
    SUPPORTED_EXTENSIONS,
    FileProblem,
    _as_text,
    _describe_read_error,
    _iter_sheets,
)

PARSERS = (
    "amount",
    "debit",
    "credit",
    "quantity",
    "unit",
    "rate",
    "date",
    "percent",
    "last_segment",
    "parent_segment",
    "strip_number",
    "map",
)
NUMERIC_PARSERS = ("amount", "debit", "credit", "quantity", "rate")
JOIN = " › "  # an upper header cell and the one below it: "Closing Balance › Debit"

# -- parsing ---------------------------------------------------------------------------------

_SEPARATORS = re.compile(r"[\s\-–—.'’]+")


def normalise(text: Any) -> str:
    """A header as presets compare it: lower case, dashes, dots, apostrophes and spaces as one space."""
    return _SEPARATORS.sub(" ", str(text or "").lower()).strip()


_SIDE = re.compile(r"(?<![A-Za-z])(dr|cr)\.?\s*$", re.I)
_CURRENCY = re.compile(r"^[A-Z]{3}\s+|[₹$€£¥]")
# A quantity is a number followed by nothing, or by a unit that starts with a letter or "/".
# Anything else ("1.2.3", "2024-04-01", "10-20") is not a quantity and is kept as written.
_QUANTITY = re.compile(r"^(\(-\)|-)?\s*([\d.,]*\d)\s*(?=$|[A-Za-z/])(.*)$")
# Digits and the separators number() reads; no exponent: "1E+5" is Excel's rounding, not a value.
_NUMERIC = re.compile(r"[\d.,'\s]*\d")
_DATE_TOKENS = re.compile(r"yyyy|yy|MMM|MM|M|dd|d")
_STRPTIME = {"yyyy": "%Y", "yy": "%y", "MMM": "%b", "MM": "%m", "M": "%m", "dd": "%d", "d": "%d"}
_NUMBERED = re.compile(r"^\s*(\S+)\s+·\s+(.+)$")


def _number(text: str) -> Optional[Decimal]:
    try:
        return Decimal(str(text).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _plain(value: Decimal) -> str:
    text = format(value.normalize(), "f")
    return "0" if text in ("-0", "") else text


def _typed(value: Any) -> bool:
    from frappe_assistant_core.plugins.core.import_opening_stock import Typed

    return isinstance(value, Typed)


def _decimal(text: str, typed: bool = False) -> Optional[Decimal]:
    """``text`` read by the opening imports' number(): commas and points as the site's number
    format says, so '1500,50' where the site writes 1.500,50 is 1500.5, not 150050. None for
    what it refuses ('12,5' where the site writes 12.5): the cell is then kept for the dry run."""
    from frappe_assistant_core.plugins.core.import_opening_stock import Typed, number

    if not _NUMERIC.fullmatch(text):
        return None
    try:
        value = number(Typed(text) if typed else text)
    except ValueError:
        return None
    return None if value is None else Decimal(repr(value))


def _signed(text: str, typed: bool = False) -> Optional[Decimal]:
    """`₹ 1,23,456.00`, `-$175,804.25`, `(1,250)`, `(-)1,250`, `8,45,000 Dr`: the number, Dr positive."""
    s = text.strip()
    sign = 1
    side = _SIDE.search(s)
    if side:
        sign = -1 if side.group(1).lower() == "cr" else 1
        s = s[: side.start()].strip()
    s = _CURRENCY.sub("", s).strip()
    if s.startswith("(-)"):
        sign, s = -sign, s[3:]
    elif s.startswith("(") and s.endswith(")"):
        sign, s = -sign, s[1:-1]
    s = s.strip()
    if s.startswith("-"):
        sign, s = -sign, s[1:].strip()
    number = _decimal(s, typed)
    return None if number is None else sign * number


def _date(text: str, formats) -> str:
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    for fmt in formats or ():
        pattern = _DATE_TOKENS.sub(lambda m: _STRPTIME[m.group(0)], fmt)
        try:
            return datetime.datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return text


def _parse_one(text: str, kind: str, spec: Dict[str, Any], typed: bool = False) -> str:
    if kind in ("amount", "debit", "credit"):
        number = _signed(text, typed)
        if number is None:
            return text
        if spec.get("sign") == "credit_positive" and _SIDE.search(text.strip()):
            number = -number  # only a value that says Dr or Cr is flipped; a plain number is as written
        if kind == "debit":
            return _plain(number) if number > 0 else ""
        if kind == "credit":
            return _plain(-number) if number < 0 else ""
        return _plain(number)
    if kind in ("quantity", "unit"):
        found = _QUANTITY.match(text)
        number = _decimal(found.group(2), typed) if found else None
        if number is None:
            return text
        if kind == "unit":
            return found.group(3).strip().lstrip("/").strip()
        return _plain(-number if found.group(1) else number)
    if kind == "rate":
        number = _signed(text.split("/", 1)[0], typed)
        return text if number is None else _plain(number)
    if kind == "date":
        return _date(text, spec.get("formats"))
    if kind == "percent":
        found = re.search(r"\d[\d.,]*", text)
        number = _decimal(found.group(0).rstrip(".,"), typed) if found else None
        return text if number is None else _plain(number)
    sep = spec.get("sep", ":")
    if kind == "last_segment":
        return text.split(sep)[-1].strip()
    if kind == "parent_segment":
        parts = text.split(sep)
        return parts[-2].strip() if len(parts) > 1 else ""
    if kind == "strip_number":
        found = _NUMBERED.match(text)
        if spec.get("keep") == "number":
            return found.group(1) if found else ""
        return found.group(2).strip() if found else text
    if kind == "map":
        values = {normalise(k): v for k, v in (spec.get("values") or {}).items()}
        return values.get(normalise(text), spec.get("default", text))
    return text


def parse(value: str, spec: Dict[str, Any]) -> str:
    """One cell, read the way a preset column says. A value no parser understands is kept as it
    was, so the dry run reports it; it never becomes 0."""
    typed = _typed(value)  # a CSV's text: '1.500' is the user's own format, not Python's 1.5
    text = (value or "").strip()
    kinds = spec.get("parse") or []
    for kind in [kinds] if isinstance(kinds, str) else kinds:
        if not text:
            break
        text = _parse_one(text, kind, spec, typed)
    if not text and "default" in spec:
        return str(spec["default"])
    return text


# -- presets ---------------------------------------------------------------------------------

PRESET_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "data", "import_presets"
)
HOOK = "assistant_import_presets"
ENGINES = {
    "opening_stock": "set_opening_stock",
    "opening_invoices": "set_opening_invoices",
    "opening_balances": "set_opening_balances",
}
GROUP_METHODS = ("empty_amounts", "indent", "known_names")
WHEN_KEYS = ("in", "not_in", "nonzero")


def _engine_fields(engine: str):
    from frappe_assistant_core.plugins.core import import_opening_balances as balances
    from frappe_assistant_core.plugins.core import import_opening_invoices as invoices
    from frappe_assistant_core.plugins.core import import_opening_stock as stock

    return {
        "opening_stock": stock.ITEM_FIELDS + stock.EXTRA_FIELDS,
        "opening_invoices": invoices.FIELDS,
        "opening_balances": balances.FIELDS,
    }[engine]


def engine_doctype(target: Dict[str, Any]) -> str:
    """The record an engine target ends up creating, for the permission check."""
    if target["engine"] == "opening_invoices":
        return f"{target.get('invoice_type')} Invoice"
    return {"opening_stock": "Stock Reconciliation", "opening_balances": "Journal Entry"}[target["engine"]]


def target_mapping(target: Dict[str, Any]) -> Dict[str, Any]:
    """Output column -> field (or candidate fields, or an engine's field), constants and inputs included."""
    mapping = dict(target.get("mapping") or target.get("columns") or {})
    for extra in (target.get("constants") or []) + (target.get("inputs") or []):
        mapping[extra["column"]] = extra["field"]
    return mapping


def _alias_groups(value, allow_empty=False) -> bool:
    if not isinstance(value, list) or (not value and not allow_empty):
        return False
    return all(
        isinstance(group, list) and group and all(isinstance(a, str) and a.strip() for a in group)
        for group in value
    )


def _is_int(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def _string_list(value: Any) -> bool:
    return isinstance(value, list) and all(isinstance(s, str) for s in value)


def _field_spec(value: Any) -> bool:
    """A field name, or a non-empty list of candidate field names."""
    if isinstance(value, str):
        return True
    return isinstance(value, list) and bool(value) and _string_list(value)


def validate_preset(preset: Any) -> Optional[str]:
    """Why ``preset`` can't be used, or None."""
    if not isinstance(preset, dict):
        return "a preset must be a JSON object"
    for key in ("id", "system", "export"):
        if not isinstance(preset.get(key), str) or not preset[key].strip():
            return f"'{key}' is required"
    signature = preset.get("signature")
    if not isinstance(signature, dict) or not _alias_groups(signature.get("required")):
        return "signature.required must be a non-empty list of alias lists"
    if signature.get("optional") is not None and not _alias_groups(signature["optional"], allow_empty=True):
        return "signature.optional must be a list of alias lists"
    if signature.get("title") is not None and not _alias_groups([signature["title"]]):
        return "signature.title must be a list of patterns"
    if signature.get("min_optional") is not None and not (
        _is_int(signature["min_optional"]) and signature["min_optional"] >= 0
    ):
        return "signature.min_optional must be a non-negative integer"
    header = preset.get("header") or {}
    if not isinstance(header, dict):
        return "header must be an object"
    if header.get("rows", 1) not in (1, 2):
        return "header.rows must be 1 or 2"
    if header.get("search_rows") is not None and not (
        _is_int(header["search_rows"]) and header["search_rows"] > 0
    ):
        return "header.search_rows must be a positive integer"

    columns = preset.get("columns")
    if not isinstance(columns, dict) or not columns:
        return "columns must be a non-empty object"
    for name, spec in columns.items():
        if not isinstance(spec, dict):
            return f"column '{name}' must be an object"
        kinds = spec.get("parse") or []
        for kind in [kinds] if isinstance(kinds, str) else kinds:
            if kind not in PARSERS:
                return f"column '{name}': unknown parse '{kind}'"
        if "map" in ([kinds] if isinstance(kinds, str) else kinds) and not isinstance(
            spec.get("values"), dict
        ):
            return f"column '{name}': parse 'map' needs 'values'"
        if spec.get("fallback") and spec["fallback"] not in columns:
            return f"column '{name}': fallback '{spec['fallback']}' is not a column"
        joined = spec.get("join")
        if joined is not None and not (
            isinstance(joined, list) and joined and all(isinstance(c, str) and c in columns for c in joined)
        ):
            return f"column '{name}': join must be a list of columns"
        if "from" in spec and not (
            isinstance(spec["from"], list) and all(isinstance(a, str) and a.strip() for a in spec["from"])
        ):
            return f"column '{name}': from must be a list of non-empty column names"
        index = spec.get("from_index")
        if index is not None and not (
            _is_int(index) or (isinstance(index, list) and all(_is_int(i) for i in index))
        ):
            return f"column '{name}': from_index must be an int or a list of ints"

    groups = preset.get("groups")
    if groups:
        if not isinstance(groups, dict):
            return "groups must be an object"
        methods = groups.get("method")
        methods = methods if isinstance(methods, list) else [methods]
        if not methods or any(m not in GROUP_METHODS for m in methods):
            return f"groups.method must be among {', '.join(GROUP_METHODS)}"
        if groups.get("column") not in columns:
            return "groups.column must be a column"
        if groups.get("carry_to") and groups["carry_to"] not in columns:
            return "groups.carry_to must be a column"
        if groups.get("names") is not None and not _string_list(groups["names"]):
            return "groups.names must be a list of strings"
    flag = preset.get("group_flag")
    if flag is not None and not (
        isinstance(flag, dict)
        and all(isinstance(flag.get(k), str) and flag[k] in columns for k in ("column", "name", "parent"))
    ):
        return "group_flag needs column, name and parent, each a column"
    skip = preset.get("skip") or {}
    if not isinstance(skip, dict):
        return "skip must be an object"
    for key in ("totals", "footer"):
        if skip.get(key) is not None and not _string_list(skip[key]):
            return f"skip.{key} must be a list of strings"
    distinct = preset.get("distinct_by")
    if distinct is not None and not isinstance(distinct, list):
        return "distinct_by must be a list"
    for key in distinct or []:
        if key not in columns:
            return f"distinct_by '{key}' is not a column"

    targets = preset.get("targets")
    if not isinstance(targets, list) or not targets:
        return "targets must be a non-empty list"
    names = set()
    for target in targets:
        if not isinstance(target, dict) or not target.get("name"):
            return "every target needs a name"
        name = target["name"]
        if name in names:
            return f"target '{name}' appears twice"
        names.add(name)
        if bool(target.get("doctype")) == bool(target.get("engine")):
            return f"target '{name}' needs exactly one of doctype or engine"
        when = target.get("when")
        if when is not None and not isinstance(when, dict):
            return f"target '{name}': when must be an object"
        if when and (when.get("column") not in columns or not any(k in when for k in WHEN_KEYS)):
            return f"target '{name}': when needs a column and one of {', '.join(WHEN_KEYS)}"
        for key in ("in", "not_in"):
            if when and key in when and not _string_list(when[key]):
                return f"target '{name}': when '{key}' must be a list of strings"
        for constant in target.get("constants") or []:
            if not all(isinstance(constant.get(k), str) for k in ("column", "field", "value")):
                return f"target '{name}': a constant needs column, field and value"
        for item in target.get("inputs") or []:
            if not all(isinstance(item.get(k), str) and item[k] for k in ("name", "column", "field", "ask")):
                return f"target '{name}': an input needs name, column, field and ask"
        extra_columns: List[str] = []
        for entry in (target.get("constants") or []) + (target.get("inputs") or []):
            if entry["column"] in columns or entry["column"] in extra_columns:
                return f"target '{name}': column '{entry['column']}' appears twice"
            extra_columns.append(entry["column"])
        base = target.get("mapping") or target.get("columns") or {}
        if not isinstance(base, dict):
            return f"target '{name}': mapping must be an object"
        for column, value in base.items():
            if target.get("engine") and not isinstance(value, str):
                return f"target '{name}': '{value}' is not a field of {target['engine']}"
            if not _field_spec(value):
                return f"target '{name}': '{column}' must map to a field or a list of fields"
        known = set(columns) | {
            e["column"] for e in (target.get("constants") or []) + (target.get("inputs") or [])
        }
        for column in target_mapping(target):
            if column not in known:
                return f"target '{name}' maps unknown column '{column}'"
        engine = target.get("engine")
        if engine:
            if engine not in ENGINES:
                return f"target '{name}': unknown engine '{engine}'"
            if engine == "opening_invoices" and target.get("invoice_type") not in ("Sales", "Purchase"):
                return f"target '{name}': invoice_type must be Sales or Purchase"
            allowed = _engine_fields(engine)
            for field in target_mapping(target).values():
                if field not in allowed:
                    return f"target '{name}': '{field}' is not a field of {engine}"
    return None


def _preset_files() -> List[str]:
    """FAC's preset files, then each hooked app's, in hook order."""
    files = sorted(glob.glob(os.path.join(PRESET_DIR, "*.json")))
    for entry in frappe.get_hooks(HOOK) or []:
        try:
            path = frappe.get_app_path(entry["app"], entry["path"])
        except Exception:
            frappe.log_error(title="Import preset hook skipped", message=repr(entry))
            continue
        files += sorted(glob.glob(os.path.join(path, "*.json"))) if os.path.isdir(path) else [path]
    return files


def _read_file(path: str) -> List[Any]:
    try:
        with open(path, encoding="utf-8") as f:
            content = json.load(f)
    except Exception as e:
        frappe.log_error(title="Import preset file skipped", message=f"{path}: {e}")
        return []
    return content if isinstance(content, list) else [content]


def load_presets() -> List[Dict[str, Any]]:
    """Every valid preset, FAC's first; a later preset with the same id replaces the earlier one.
    Read once per request."""
    cached = getattr(frappe.local, "fac_import_presets", None)
    if cached is not None:
        return cached
    by_id: Dict[str, Dict[str, Any]] = {}
    for path in _preset_files():
        for preset in _read_file(path):
            try:
                problem = validate_preset(preset)
            except Exception as e:
                # A shape validate_preset does not expect is this preset's problem only.
                problem = f"{type(e).__name__}: {e}"
            if problem:
                ident = preset.get("id") if isinstance(preset, dict) else None
                frappe.log_error(title="Import preset skipped", message=f"{path}: {ident}: {problem}")
                continue
            by_id[preset["id"]] = preset
    frappe.local.fac_import_presets = list(by_id.values())
    return frappe.local.fac_import_presets


def clear_cache() -> None:
    frappe.local.fac_import_presets = None


# -- reading ---------------------------------------------------------------------------------

SEARCH_ROWS = 20
MIN_HEAD_ROWS = 22  # SEARCH_ROWS plus the second row of a two-row header, and one spare
MAX_PRESET_ROWS = 200_000  # a matched sheet is read whole; past this, the normal flow streams it
DROP_REASONS = ("title", "blank", "repeated_header", "totals", "footer", "groups", "merged", "no_target")


@dataclass
class Sheet:
    name: Optional[str]
    rows: List[List[str]]
    indents: Optional[List[List[int]]] = None  # xlsx cell indent per cell, when a preset needs it


def _extension(file_doc) -> str:
    return (file_doc.get_extension()[1] or "").lstrip(".").lower()


def _content(file_doc) -> bytes:
    if _extension(file_doc) not in SUPPORTED_EXTENSIONS:
        raise FileProblem(_("Only .csv, .xlsx or .xls files can be imported."))
    if (file_doc.file_size or 0) > MAX_FILE_SIZE:
        raise FileProblem(_("The file is larger than 50 MB."))
    content = file_doc.get_content()
    return content.encode("utf-8") if isinstance(content, str) else content


def _csv_rows(content: bytes) -> Iterator[List[str]]:
    """A CSV's rows with cells as written, one at a time. Frappe's reader strips every cell, which
    loses the leading spaces that mark groups by indent. QuickBooks Desktop writes cp1252."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return (list(row) for row in csv.reader(io.StringIO(text), dialect))


def _sheet_rows(file_doc) -> Iterator[Tuple[Optional[str], Iterator[List[str]]]]:
    """(sheet name, rows as text) for each sheet, lazily: a sheet is read only as far as it is
    consumed, and the next one only after it. A CSV is one sheet named None."""
    content = _content(file_doc)
    if _extension(file_doc) == "csv":
        from frappe_assistant_core.plugins.core.import_opening_stock import Typed

        yield None, ([Typed(cell) for cell in row] for row in _csv_rows(content))
        return
    for name, rows in _iter_sheets(content, _extension(file_doc)):
        yield name, ([_as_text(v) for v in row] for row in rows)


def _reading(read):
    """``read()``, with any failure to read the file as a FileProblem."""
    try:
        return read()
    except FileProblem:
        raise
    except Exception as e:
        raise FileProblem(_("Could not read the file: {0}").format(_describe_read_error(e)))


def head_rows(preset_list: Optional[List[Dict[str, Any]]] = None) -> int:
    """How many rows matching needs from the top of a sheet: the largest header search, plus two
    for a two-row header, and never fewer than 22."""
    loaded = load_presets() if preset_list is None else preset_list
    deepest = max(
        ((p.get("header") or {}).get("search_rows", SEARCH_ROWS) for p in loaded), default=SEARCH_ROWS
    )
    return max(deepest + 2, MIN_HEAD_ROWS)


def read_heads(file_doc, rows: Optional[int] = None) -> List[Sheet]:
    """The first ``rows`` rows (default: head_rows()) of every sheet, as text. Enough to match a
    preset; the rest of each sheet is never read."""
    limit = head_rows() if rows is None else rows
    return _reading(
        lambda: [Sheet(name, list(itertools.islice(cells, limit))) for name, cells in _sheet_rows(file_doc)]
    )


def read_sheet(file_doc, sheet_name: Optional[str]) -> Sheet:
    """One sheet in full (for a CSV, sheet_name is None), up to MAX_PRESET_ROWS rows."""

    def read():
        for name, cells in _sheet_rows(file_doc):
            if name != sheet_name and sheet_name is not None:
                continue
            rows = list(itertools.islice(cells, MAX_PRESET_ROWS + 1))
            if len(rows) > MAX_PRESET_ROWS:
                raise FileProblem(
                    _(
                        "This sheet has more than {0} rows. Split it into smaller files, or use "
                        "start_import_session."
                    ).format(f"{MAX_PRESET_ROWS:,}")
                )
            return Sheet(name, rows)
        raise FileProblem(_("Sheet '{0}' not found.").format(sheet_name))

    return _reading(read)


def read_sheets(file_doc) -> List[Sheet]:
    """Every sheet of the file, every row as text, title rows and all. The tool does not use it:
    it reads read_heads, then read_sheet for the matched sheet only."""
    return _reading(lambda: [Sheet(name, list(cells)) for name, cells in _sheet_rows(file_doc)])


def read_indents(file_doc, sheet_name: Optional[str]) -> Optional[List[List[int]]]:
    """Each cell's indent in an .xlsx sheet. Read-only mode doesn't load cell styles, so the
    workbook is opened in full; only done when a matched preset finds groups by indent."""
    if _extension(file_doc) != "xlsx":
        return None

    workbook = load_workbook(io.BytesIO(_content(file_doc)), data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        return [[int(cell.alignment.indent or 0) for cell in row] for row in worksheet.iter_rows(min_row=1)]
    finally:
        workbook.close()


def uses_indent(preset: Dict[str, Any]) -> bool:
    methods = (preset.get("groups") or {}).get("method") or []
    return "indent" in (methods if isinstance(methods, list) else [methods])


# -- matching --------------------------------------------------------------------------------


@dataclass
class Match:
    preset: Dict[str, Any]
    header_row: int  # 0-based index of the (first) header row
    headers: List[str]
    score: float
    hits: int


def _cell(row: List[str], index: Optional[int]) -> str:
    return row[index] if index is not None and index < len(row) else ""


def _header_cells(rows: List[List[str]], index: int, depth: int) -> List[str]:
    upper = rows[index]
    if depth == 1:
        return [c.strip() for c in upper]
    lower = rows[index + 1] if index + 1 < len(rows) else []
    out, carried = [], ""
    for i in range(max(len(upper), len(lower))):
        above, below = _cell(upper, i).strip(), _cell(lower, i).strip()
        # A merged cell reads as blanks to its right: carry it over the cells that have a header below.
        carried = above or (carried if below else "")
        out.append(f"{carried}{JOIN}{below}" if carried and below else (carried or below))
    return out


def _found(group: List[str], names) -> bool:
    return any(normalise(alias) in names for alias in group)


def _like(text: str, patterns: Iterable[str]) -> bool:
    value = normalise(text)
    return any(fnmatch.fnmatchcase(value, normalise(p)) for p in patterns)


def _title_ok(signature: Dict[str, Any], rows: List[List[str]], index: int) -> bool:
    patterns = signature.get("title")
    if not patterns:
        return True
    return any(_like(c, patterns) for row in rows[:index] for c in row if c.strip())


def match_sheet(rows: List[List[str]], presets: List[Dict[str, Any]]) -> List[Match]:
    """Every preset whose signature this sheet has, best first; load order breaks ties."""
    found: List[Match] = []
    for preset in presets:
        header = preset.get("header") or {}
        depth = header.get("rows", 1)
        signature = preset["signature"]
        optional = signature.get("optional") or []
        min_optional = signature.get("min_optional") or 0
        has_title = 1 if signature.get("title") else 0
        best: Optional[Match] = None
        for index in range(min(header.get("search_rows", SEARCH_ROWS), len(rows))):
            headers = _header_cells(rows, index, depth)
            names = {normalise(h) for h in headers if h.strip()}
            if not all(_found(g, names) for g in signature["required"]):
                continue
            optional_hits = sum(_found(g, names) for g in optional)
            # A generic signature ("Name, Group") also needs some of the system's own columns.
            if optional_hits < min_optional:
                continue
            if not _title_ok(signature, rows, index):
                continue
            hits = len(signature["required"]) + optional_hits + has_title
            total = len(signature["required"]) + len(optional) + has_title
            # The header is the row with the most groups found; a tie goes to the earlier row.
            if best is None or hits > best.hits:
                best = Match(preset, index, headers, hits / total, hits)
        if best:
            found.append(best)
    found.sort(key=lambda m: -m.score)  # stable: equal scores keep load order
    return found


def missing_groups(rows: List[List[str]], preset: Dict[str, Any]) -> List[List[str]]:
    """The required header groups no row in the search window has, for a preset the user named."""
    header = preset.get("header") or {}
    limit = min(header.get("search_rows", SEARCH_ROWS), len(rows))
    names = {
        normalise(h)
        for index in range(limit)
        for h in _header_cells(rows, index, header.get("rows", 1))
        if h.strip()
    }
    return [g for g in preset["signature"]["required"] if not _found(g, names)]


def why_not(rows: List[List[str]], preset: Dict[str, Any]) -> Dict[str, Any]:
    """Why a preset the user named does not match: ``missing_headers`` (the required groups not
    found), and ``missing_title`` (the title patterns, when none is above any row that has the
    required groups) or ``needs_optional`` (when no such row has min_optional optional groups)."""
    out: Dict[str, Any] = {"missing_headers": missing_groups(rows, preset)}
    header = preset.get("header") or {}
    signature = preset["signature"]
    optional = signature.get("optional") or []
    min_optional = signature.get("min_optional") or 0
    candidates, enough = [], []
    for index in range(min(header.get("search_rows", SEARCH_ROWS), len(rows))):
        names = {normalise(h) for h in _header_cells(rows, index, header.get("rows", 1)) if h.strip()}
        if not all(_found(g, names) for g in signature["required"]):
            continue
        candidates.append(index)
        if sum(_found(g, names) for g in optional) >= min_optional:
            enough.append(index)
    if not candidates:
        return out
    if not enough:
        out["needs_optional"] = {"min": min_optional, "groups": optional}
    if signature.get("title") and not any(_title_ok(signature, rows, i) for i in enough or candidates):
        out["missing_title"] = signature["title"]
    return out


# -- cleaning --------------------------------------------------------------------------------


@dataclass
class TargetRows:
    target: Dict[str, Any]
    header: List[str]
    rows: List[List[str]] = field(default_factory=list)


@dataclass
class Cleaned:
    targets: List[TargetRows]
    dropped: Dict[str, int]
    no_target_values: Dict[str, List[str]]
    rows_read: int


def _positions(preset: Dict[str, Any], headers: List[str]) -> Dict[str, List[int]]:
    """Output column -> the source positions to read it from (the first filled one wins)."""
    names = [normalise(h) for h in headers]
    out: Dict[str, List[int]] = {}
    for column, spec in preset["columns"].items():
        for alias in spec.get("from") or []:
            if normalise(alias) in names:
                out[column] = [names.index(normalise(alias))]
                break
        else:
            index = spec.get("from_index")
            if index is not None:
                out[column] = [i for i in (index if isinstance(index, list) else [index]) if i < len(headers)]
    return out


def _read(row: List[str], positions: List[int]) -> str:
    for i in positions:
        if _cell(row, i).strip():
            return _cell(row, i)
    return ""


def _level(sheet: Sheet, r: int, positions: List[int]) -> int:
    col = next((i for i in positions if _cell(sheet.rows[r], i).strip()), positions[0] if positions else 0)
    text = _cell(sheet.rows[r], col)
    indent = 0
    if sheet.indents and r < len(sheet.indents) and col < len(sheet.indents[r]):
        indent = sheet.indents[r][col]
    return indent * 100 + len(text) - len(text.lstrip(" "))


def _group_rows(sheet: Sheet, kept: List[int], preset: Dict[str, Any], positions) -> set:
    spec = preset.get("groups")
    if not spec:
        return set()
    methods = spec["method"] if isinstance(spec["method"], list) else [spec["method"]]
    at = positions.get(spec["column"], [])
    numeric = [
        positions[c]
        for c, s in preset["columns"].items()
        if c in positions
        and set([s.get("parse")] if isinstance(s.get("parse"), str) else s.get("parse") or [])
        & set(NUMERIC_PARSERS)
    ]
    names = {normalise(n) for n in spec.get("names") or []}
    out = set()
    for i, r in enumerate(kept):
        row = sheet.rows[r]
        name = _read(row, at)
        if not name.strip():
            continue
        if "empty_amounts" in methods and numeric and not any(_read(row, p).strip() for p in numeric):
            out.add(r)
        elif "known_names" in methods and normalise(name) in names:
            out.add(r)
        elif (
            "indent" in methods
            and i + 1 < len(kept)
            and _level(sheet, kept[i + 1], at) > _level(sheet, r, at)
        ):
            out.add(r)
    return out


def _when(when: Optional[Dict[str, Any]], values: Dict[str, str]) -> bool:
    if not when:
        return True
    value = values.get(when["column"], "")
    if "in" in when:
        return normalise(value) in {normalise(v) for v in when["in"]}
    if "not_in" in when:
        return normalise(value) not in {normalise(v) for v in when["not_in"]}
    # Only blank and a number equal to zero fail: text that isn't a number ("n/a") goes on to the
    # dry run, which reports it, instead of being dropped here as fitting no target.
    if not value.strip():
        return False
    number = _number(value)
    return number is None or not number.is_zero()


def clean(sheet: Sheet, match: Match) -> Cleaned:
    """The sheet's data rows, read the preset's way and split by target: title, blank, repeated
    header, total, footer and group rows out, and counted."""
    preset = match.preset
    depth = (preset.get("header") or {}).get("rows", 1)
    skip = preset.get("skip") or {}
    positions = _positions(preset, match.headers)
    header_row = [normalise(c) for c in sheet.rows[match.header_row]]
    dropped = dict.fromkeys(DROP_REASONS, 0)
    dropped["title"] = match.header_row

    kept: List[int] = []
    totals: set = set()
    first_data = match.header_row + depth
    for r in range(first_data, len(sheet.rows)):
        row = sheet.rows[r]
        filled = [c for c in row if c.strip()]
        if not filled:
            dropped["blank"] += 1
        elif [normalise(c) for c in row[: len(header_row)]] + [""] * (
            len(header_row) - len(row)
        ) == header_row:
            dropped["repeated_header"] += 1
        elif _like(filled[0], skip.get("totals") or []):
            dropped["totals"] += 1
            totals.add(r)
        elif len(filled) == 1 and _like(filled[0], skip.get("footer") or []):
            dropped["footer"] += 1
        else:
            kept.append(r)

    groups = _group_rows(sheet, kept, preset, positions)
    group_spec = preset.get("groups") or {}
    carry = group_spec.get("carry_to")
    current_group = ""
    built: List[Dict[str, str]] = []
    for r in sorted(totals.union(kept)):
        row = sheet.rows[r]
        if r in totals:
            current_group = ""  # "Total for X" closes X
            continue
        if r in groups:
            dropped["groups"] += 1
            current_group = _read(row, positions.get(group_spec["column"], [])).strip()
            continue
        values = {}
        for column, spec in preset["columns"].items():
            if column == carry:
                values[column] = current_group
            else:
                values[column] = parse(_read(row, positions.get(column, [])), spec)
        for column, spec in preset["columns"].items():
            if spec.get("fallback") and not values[column]:
                values[column] = values.get(spec["fallback"], "")
        for column, spec in preset["columns"].items():
            if spec.get("join"):
                parts = [values[c] for c in spec["join"] if values[c].strip()]
                values[column] = spec.get("sep", ":").join(parts)
        built.append(values)

    if preset.get("distinct_by"):
        seen, unique = set(), []
        for values in built:
            key = tuple(values[c] for c in preset["distinct_by"])
            if key in seen:
                dropped["merged"] += 1
                continue
            seen.add(key)
            unique.append(values)
        built = unique

    flag = preset.get("group_flag")
    if flag:
        parents = {normalise(v[flag["parent"]]) for v in built if v[flag["parent"]].strip()}
        for values in built:
            name = values[flag["name"]]
            values[flag["column"]] = "1" if name.strip() and normalise(name) in parents else "0"

    columns = list(preset["columns"])
    targets = [TargetRows(t, list(columns)) for t in preset["targets"]]
    split_on = list(dict.fromkeys(t["when"]["column"] for t in preset["targets"] if t.get("when")))
    no_target: Dict[str, List[str]] = {}
    for values in built:
        for target in targets:
            if _when(target.target.get("when"), values):
                target.rows.append([values[c] for c in columns])
                break
        else:
            dropped["no_target"] += 1
            for column in split_on:
                seen_values = no_target.setdefault(column, [])
                if values[column] not in seen_values and len(seen_values) < 20:
                    seen_values.append(values[column])
    return Cleaned(targets, dropped, no_target, rows_read=max(len(sheet.rows) - first_data, 0))


def missing_inputs(target: Dict[str, Any], values: Dict[str, Any]) -> List[Dict[str, Any]]:
    """The inputs a target needs that the user hasn't given yet."""
    return [i for i in target.get("inputs") or [] if not str((values or {}).get(i["name"]) or "").strip()]


def finish(rows: TargetRows, values: Dict[str, Any]) -> TargetRows:
    """The target's rows with its constant and input columns added."""
    extra: List[Tuple[str, str]] = [(c["column"], c["value"]) for c in rows.target.get("constants") or []]
    extra += [(i["column"], str(values[i["name"]]).strip()) for i in rows.target.get("inputs") or []]
    return TargetRows(
        rows.target,
        rows.header + [column for column, _value in extra],
        [row + [value for _column, value in extra] for row in rows.rows],
    )
